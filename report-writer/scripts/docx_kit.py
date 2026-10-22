# -*- coding: utf-8 -*-
"""docx_kit —— 往既有 Word 模板里“插入”内容的工具集。

用法：写一份本次报告专用的 build.py，import 本模块，然后按文档顺序调用这些
helper 把内容塞进模板的可编写区域。不要用 python-docx 重建文档——模板的框、
页眉页脚、成绩表、样式定义都在里面，重建必丢。

三块内容要分开对待：

| 区域 | 能不能写 | 怎么写 |
| --- | --- | --- |
| **封面 / 表头字段**（姓名、学号、题目…） | 能，但**只能填值** | `K.cover_fields()` 勘察 → 问用户 → `K.fill_cover()` |
| **正文编写区** | 能 | `K.Cursor` + `K.Target` |
| **内嵌表格**（题目表、信息表、评分表） | **不能整格清空** | 用 `Target.after('锚点')` 或 `Cursor(..., clear=False)` |

    import sys; sys.path.insert(0, r'<skill>/scripts')
    from docx import Document
    import docx_kit as K

    d = Document('work.docx')            # 由 word_com.ps1 从 .doc 转来

    # ① 封面：只填用户给了值的字段，标签/下划线/图片一字不动
    for e in K.cover_fields(d):
        print(e['where'], e['label'], repr(e['value']))
    K.fill_cover(d, {'姓名': '示例姓名', '学号': '00000000'})

    # ② 正文：定位那个“框”，往里写（不要清空含内嵌表格的格子）
    cell = K.editable_cell(K.find_frame(d)[0], row=3)
    K.clear_cell(cell, keep=1)           # 保留模板原有的标题段

    caps = K.Captions()
    caps.declare('fig_topo', '图')        # 第一遍：按最终出现顺序登记编号
    caps.declare('tab_host', '表')

    K.put(cell, K.heading(1, '1. 网络拓扑搭建'))
    K.put(cell, K.body(caps, '拓扑见<<fig_topo>>。'))
    K.put(cell, K.image(cell, 'Pictures/topo.png', width_pt=400))
    K.put(cell, K.caption(caps, 'fig_topo', '实验网络拓扑'))
    K.put(cell, K.table_block(caps, 'tab_host', '主机地址规划', rows, widths))
    K.finish_cell(cell)                  # 单元格不能以表格结尾
    d.save('out.docx')

之后用 word_com.ps1 更新域并另存为 .doc / 导出 PDF；用 page_fill.py 量页面利用率。
"""

import copy
import hashlib
import os
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

# --------------------------------------------------------------------------
# 底层：元素与 run
# --------------------------------------------------------------------------


def E(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn('w:' + k), v)
    return el


def rpr(fonts=None, eastasia=None, bold=False, color=None, sz=None,
        noproof=False, superscript=False, underline=False, hint='eastAsia'):
    """构造 <w:rPr>。子元素顺序必须符合 OOXML schema，写错 Word 会判定文件损坏：
    rFonts → b → bCs → noProof → color → sz → u → vertAlign。

    superscript=True 生成上标（文献引用 [1] 要的右上角标）。
    underline=True 生成下划线（封面字段填值常用）。"""
    r = OxmlElement('w:rPr')
    if fonts or eastasia or hint:
        f = OxmlElement('w:rFonts')
        if fonts:
            f.set(qn('w:ascii'), fonts)
            f.set(qn('w:hAnsi'), fonts)
        if eastasia:
            f.set(qn('w:eastAsia'), eastasia)
        if hint:
            f.set(qn('w:hint'), hint)
        r.append(f)
    if bold:
        r.append(OxmlElement('w:b'))
        r.append(OxmlElement('w:bCs'))
    if noproof:
        r.append(OxmlElement('w:noProof'))
    if color:
        r.append(E('w:color', val=color))
    if sz:
        r.append(E('w:sz', val=str(sz)))
    if underline:
        r.append(E('w:u', val='single'))
    if superscript:
        r.append(E('w:vertAlign', val='superscript'))
    return r if len(r) else None


def run(text, prop=None, preserve=False):
    r = OxmlElement('w:r')
    if prop is not None and len(prop):
        r.append(copy.deepcopy(prop))
    t = OxmlElement('w:t')
    if preserve or (text != text.strip() and text.strip()):
        t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return r


def field(instr, cached):
    """构造一个 Word 域。cached 是域结果——务必填对，因为域没被刷新时
    读到的就是它，而很多人打开文档根本不会按 F9。"""
    out = []
    r = OxmlElement('w:r')
    r.append(E('w:fldChar', fldCharType='begin'))
    out.append(r)
    r = OxmlElement('w:r')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = instr
    r.append(it)
    out.append(r)
    r = OxmlElement('w:r')
    r.append(E('w:fldChar', fldCharType='separate'))
    out.append(r)
    r = OxmlElement('w:r')
    r.append(rpr(noproof=True))
    t = OxmlElement('w:t')
    t.text = cached
    r.append(t)
    out.append(r)
    r = OxmlElement('w:r')
    r.append(E('w:fldChar', fldCharType='end'))
    out.append(r)
    return out


def new_p(style=None, keep_next=False, keep_lines=False, center=False,
          indent=False, spacing=None, tiny=False, shade=None):
    """新建 <w:p>。pPr 子元素同样有顺序要求：
    pStyle → keepNext → keepLines → shd → snapToGrid → spacing → ind → jc → rPr。

    shade 是灰底填充色（如 'F2F2F2'），代码块用它。"""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    if style:
        pPr.append(E('w:pStyle', val=style))
    if keep_next:
        pPr.append(OxmlElement('w:keepNext'))
    if keep_lines:
        pPr.append(OxmlElement('w:keepLines'))
    if shade:
        # shd 必须排在 pBdr 之后、snapToGrid 之前，顺序错了 Word 报文件损坏
        pPr.append(E('w:shd', val='clear', color='auto', fill=shade))
    pPr.append(E('w:snapToGrid', val='0'))
    if spacing:
        pPr.append(E('w:spacing', before=str(spacing[0]), after=str(spacing[1])))
    if indent:
        pPr.append(E('w:ind', firstLineChars='200', firstLine='420'))
    if center:
        pPr.append(E('w:jc', val='center'))
    if tiny:
        # 1pt 的段落标记：用于“单元格不能以表格结尾”这种占位，不占一整行
        r = OxmlElement('w:rPr')
        r.append(E('w:sz', val='2'))
        r.append(E('w:szCs', val='2'))
        pPr.append(r)
    p.append(pPr)
    return p


# --------------------------------------------------------------------------
# 结构探查：找到“框”和可编写区域
# --------------------------------------------------------------------------


def _row_cells(tr):
    """该行【去重后】的 tc 元素。python-docx 的 row.cells 会把横向合并的格子
    重复返回，直接 len(row.cells) 会把 2 列的表算成 9 列——判断列数必须用这个。"""
    return tr.findall(qn('w:tc'))


def _tc_text(tc):
    return ''.join(t.text or '' for t in tc.iter(qn('w:t'))).strip()


def _effective_cols(tbl_el):
    """各行的有效格数集合。只有所有行都只有 1 格，才是真正的单列表。"""
    return set(len(_row_cells(tr)) for tr in tbl_el.findall(qn('w:tr')))


RUBRIC_HINTS = ('评价', '赋分', '得分', '评分', '考核', '成绩', '分值',
                '权重', '评分标准', '评阅', '分项', '教师', '教学目标')


def looks_like_rubric(tbl_el):
    """是不是评分表。评分表是**绝对不能往里写正文**的地方——
    它长得又像"有很多空格的表格"，是自动识别最容易踩的坑。"""
    trs = tbl_el.findall(qn('w:tr'))
    head = ' '.join(_tc_text(tc) for tr in trs[:2] for tc in _row_cells(tr))
    return sum(1 for h in RUBRIC_HINTS if h in head) >= 2


def _is_bordered(tbl_el):
    pr = tbl_el.find(qn('w:tblPr'))
    return pr is not None and pr.find(qn('w:tblBorders')) is not None


def row_looks_like_rubric(tbl_el, row_idx, look_up=3):
    """**按行**判断是不是评分行。

    为什么不能整表判断：有的模板把"填写栏"和"教师评分表"做在同一张表里
    （前半是问题介绍/算法实现，后半是分项打分）。整表判会误伤前半。
    所以看目标行**及其上方几行**——评分表的分项标题常在上一行。
    """
    trs = tbl_el.findall(qn('w:tr'))
    seg = trs[max(0, row_idx - look_up): row_idx + 1]
    txt = ' '.join(_tc_text(tc) for tr in seg for tc in _row_cells(tr))
    return sum(1 for h in RUBRIC_HINTS if h in txt) >= 2


def tables(doc):
    """**顶层**表格元素列表（按文档顺序）。

    注意与 `all_tables()` 的区别：`Target.cell(table=N)` 里的 N 数的是这个
    顶层序号；而结构校验要用 `all_tables()`——它连内嵌表格一起管。
    """
    return doc.element.body.findall(qn('w:tbl'))


def iter_tables(root, prefix=''):
    """递归收集表格，产出 `(标签, 表元素)`。

    为什么必须递归：模板里除了最外面那个框，多半还有**内嵌表格**——题目表格、
    学生信息表、评分表，全都长在某个单元格里。不递归的话它们不在结构指纹里，
    **被改坏了也不会报出来**。

    标签形如 `TABLE[1]`（顶层）、`TABLE[1] row3 col0 内嵌表[0]`（嵌在哪一格）。
    """
    out = []
    for i, tbl in enumerate(root.findall(qn('w:tbl'))):
        label = 'TABLE[%d]' % i if not prefix else '%s 内嵌表[%d]' % (prefix, i)
        out.append((label, tbl))
        for ri, tr in enumerate(tbl.findall(qn('w:tr'))):
            for ci, tc in enumerate(_row_cells(tr)):
                out.extend(iter_tables(tc, '%s row%d col%d' % (label, ri, ci)))
    return out


def all_tables(doc):
    """文档里**所有**表格元素（含内嵌），按出现顺序。"""
    return [el for _, el in iter_tables(doc.element.body)]


def nested_tables(doc):
    """只挑出**内嵌**表格（不在顶层的那批）：`[(标签, 元素), ...]`。

    它们是模板里的题目表格/信息表/评分表，任何情况下都不要搬动或清空。
    """
    return [(lab, el) for lab, el in iter_tables(doc.element.body)
            if '内嵌表' in lab]


def find_frame(doc, min_rows=2):
    """只返回**单列长表**，找不到就返回空列表。

    为什么收得这么紧：模板里"留给作者填写的区域"至少有四种形态——
    ① 单列带边框长表 ② 单列无边框长表 ③ 多列表单（每行一个栏目）
    ④ 纯段落结构。除了第①种，其余都无法靠启发式可靠识别，而**认错的代价
    是静默地把正文写进评分表**。

    所以这里只认唯一能安全自动识别的那一种；其余形态请用 `survey(doc)` 勘察
    后**明确指定**目标（见 `Target`）。宁可返回空让你去指定，也不要猜一个出来。
    """
    out = []
    for tbl_el in tables(doc):
        trs = tbl_el.findall(qn('w:tr'))
        if len(trs) < min_rows:
            continue
        if _effective_cols(tbl_el) != {1}:        # 必须每行都只有 1 格
            continue
        if looks_like_rubric(tbl_el):
            continue
        chars = sum(len(_tc_text(tc)) for tr in trs for tc in _row_cells(tr))
        out.append((-chars, tbl_el))
    out.sort(key=lambda x: x[0])                  # 内容多的排前面
    return [Table(el, doc._body) for _, el in out]


COVER_HINTS = ('姓名', '学号', '班级', '日期', '题目', '报告名称', '学生',
               '实验名称', '实验地点', '组员', '成员', '指导教师', '专业', '课程',
               '项目名称', '项目成员', '年级', '学院', '单位', '完成日期')


def _norm(s):
    """去掉所有空白（含全角空格）再比对 —— 模板爱写成「题　　目」「班　　级」，
    直接做子串匹配会因为中间那个全角空格而漏掉。"""
    return re.sub(r'[\s　]+', '', s or '')


def _rubric_start(tbl_el):
    """评分表从第几行开始（那一行本身就是评分表头）。不是评分表返回 None。

    比"逐行判"更稳：一旦认定从第 N 行起是评分区，**之后所有行都不建议写**，
    免得评分表里那些空着的"得分"格被当成填写区。
    """
    for ri, tr in enumerate(tbl_el.findall(qn('w:tr'))):
        txt = ' '.join(_tc_text(tc) for tc in _row_cells(tr))
        if sum(1 for h in RUBRIC_HINTS if h in txt) >= 2:
            return ri
    return None

# 模板常常在填写区里留一句占位说明，而不是留空
PLACEHOLDER_HINTS = ('注：', '此栏目', '请在此', '不少于', 'XXX', 'xxx',
                     '……', '正文字数', '（删', '可参考')


RUBRIC_LABELS = ('评语', '评阅', '评分', '得分', '成绩', '考核', '评价')


def _is_rubric_label(cells):
    """行标签本身就是评分类（「七、教师评语：」）—— 单个词就够排除了。"""
    return bool(cells) and any(h in _norm(_tc_text(cells[0])) for h in RUBRIC_LABELS)


def _is_cover_row(cells):
    """封面信息行（姓名/学号这类）不是正文填写区。"""
    return bool(cells) and any(h in _norm(_tc_text(cells[0])) for h in COVER_HINTS)


def _writable(text, threshold):
    """空着、只留一句占位说明、或只留了一行小节标题 —— 都算"留给你写的"。

    最后那条很关键：模板常常写成「四、基础实验步骤与结果：」独占一格，正文写在这格
    的下面。这种格子有十几个字，光按"空"判会漏掉。"""
    n = len(text)
    if n <= threshold:
        return True
    if any(h in text for h in PLACEHOLDER_HINTS) and n <= 80:
        return True
    return n <= 24 and text.rstrip().endswith(('：', ':'))


def plan(doc, empty_threshold=10, max_candidates=8):
    """给出**写入方案建议**：可直接复制的 Target、要避开的区域、下一步代码。

    `survey()` 摆事实，`plan()` 在此基础上给建议——但仍然只是建议：
    真正写入前必须过 `Cursor.check()`，写完必须过 `finish()` 比对。
    """
    L = ['模板写入方案', '-' * 58]
    tbls = tables(doc)
    shape = None
    cands = []          # (说明, 代码)
    cover = cover_fields(doc)   # 封面/表头字段（结构化）
    skipped = []        # 评分区
    nested = nested_tables(doc)

    for ti, el in enumerate(tbls):
        trs = el.findall(qn('w:tr'))
        cols = _effective_cols(el)
        rs = _rubric_start(el)                      # 第几行起是评分区
        if rs is not None:
            skipped.append('TABLE[%d]（第 %d 行起）' % (ti, rs))
        for ri, tr in enumerate(trs):
            if rs is not None and ri >= rs:
                continue                            # 评分区一律不建议
            cells = _row_cells(tr)
            if _is_cover_row(cells):
                continue        # 封面行不进正文候选，单独在下面列（cover_fields）
            if _is_rubric_label(cells):
                continue
            if not _writable(_tc_text(cells[0]), 0) and cols != {1} and len(cells) > 1:
                pass
            if cols == {1}:
                for ci, tc in enumerate(cells):
                    txt = _tc_text(tc)
                    if not _writable(txt, empty_threshold):
                        continue
                    shape = shape or '① 单列长表'
                    cands.append((
                        'TABLE[%d] row%d 「%s」 现有 %d 字' % (ti, ri, txt[:16], len(txt)),
                        'Target.cell(table=%d, row=%d, col=%d)' % (ti, ri, ci)))
            else:
                for ci, tc in enumerate(cells):
                    if ci == 0 or not _writable(_tc_text(tc), empty_threshold):
                        continue
                    shape = shape or '② 多列表单'
                    cands.append((
                        'TABLE[%d] row%d 栏目「%s」' % (ti, ri, _tc_text(cells[0])[:16]),
                        'Target.cell(table=%d, row=%d, col=%d)' % (ti, ri, ci)))

    if not cands:
        # 兜底一：单列框的每行都被模板预填了说明文字 —— 这时"要写的位置"往往是
        # 某一行之内（在其后追加）。挑最短的几行作为推测，交给人工定夺。
        for ti, el in enumerate(tbls):
            if _effective_cols(el) != {1}:
                continue
            rs = _rubric_start(el)
            rows = [(ri, _tc_text(_row_cells(tr)[0]))
                    for ri, tr in enumerate(el.findall(qn('w:tr')))
                    if (rs is None or ri < rs)]
            rows = [(ri, t) for ri, t in rows
                    if t and not any(h in _norm(t) for h in RUBRIC_LABELS)]
            picked = sorted(rows, key=lambda x: len(x[1]))[:3]
            for ri, t in picked:
                cands.append(('[推测] TABLE[%d] row%d 「%s」' % (ti, ri, t[:20]),
                              'Target.cell(table=%d, row=%d, col=0)' % (ti, ri)))
            if picked:
                shape = '③ 行内预填、需在行内追加（需人工确认）'
                break

    if not cands:
        # 兜底二：**正文顶层**的段落锚点。注意只扫 body 的直接子段落 ——
        # 用 iter() 会把表格单元格里的段落也捞进来，那些已经在上面处理过了。
        for el in doc.element.body.findall(qn('w:p')):
            txt = ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()
            if not (0 < len(txt) <= 40):
                continue
            is_head = bool(re.match(r'(\d+[\.、]|[一二三四五六七八九十]+、|\d+\.\d+)', txt))
            if not (txt.endswith(('：', ':')) or is_head
                    or any(h in txt for h in PLACEHOLDER_HINTS)):
                continue
            if any(h in _norm(txt) for h in COVER_HINTS):
                continue        # 段落式封面字段同上，cover_fields 里已收录
            cands.append(('段落锚点「%s」' % txt, "Target.after('%s')" % txt))
        if cands:
            shape = '④ 纯段落结构'

    L.append('形态判定：%s' % (shape or '③ 正文区与评分表混排（需人工判断）'))
    L.append('')
    L.append('写入目标候选（**仍需你确认**）：')
    if cands:
        for desc, code in cands[:max_candidates]:
            L.append('    %-42s →  K.%s' % (desc, code))
        if len(cands) > max_candidates:
            L.append('    ...（还有 %d 个，见 survey()）' % (len(cands) - max_candidates))
    else:
        L.append('    （没有明显候选，请人工阅读 survey() 的输出后指定）')
    L.append('')

    if cover:
        pending = [f for f in cover if not f['value']]
        L.append('封面/表头字段：共 %d 个，其中 %d 个还空着（**问用户**，别猜）：'
                 % (len(cover), len(pending)))
        for f in cover[:10]:
            L.append('    [%s] %-12s 值「%s」  %s' % (
                '待填' if not f['value'] else '已填',
                f['label'], f['value'][:18], f['where']))
        if len(cover) > 10:
            L.append('    ...（共 %d 个）' % len(cover))
        L.append('    → 填用 K.fill_cover(doc, {...})：只改值，不动标签、不清空、不删框。')
        L.append('      空着的字段若用户选择手填，就**一个字都别动封面**。')
        L.append('')
    if nested:
        L.append('[!] 模板里有 %d 张**内嵌表格**（多半是题目表/信息表/评分表）：' % len(nested))
        for lab, el in nested[:4]:
            L.append('    %s  %d行  「%s」' % (
                lab, len(el.findall(qn('w:tr'))), _tbl_first_text(el)[:20]))
        L.append('    → 它们已被纳入结构指纹，被改动会在 finish() 里报出来。')
        L.append('      往它们所在的格子写内容：Target.after("锚点") 或 Cursor(..., clear=False)，')
        L.append('      **不要用 Target.cell() 清空整格**——清空会把格内原有的表格顶到最前。')
        L.append('')
    if skipped:
        L.append('[!] 已排除评分区：%s' % '、'.join(skipped))
        L.append('')

    L.append('下一步：')
    if cands:
        L.append('    target = K.%s' % cands[0][1])
        L.append('    cur = K.Cursor(d, target, keep_paras=1)')
        L.append('    assert not cur.check(), cur.check()')
        L.append('    ... cur.put(...) ...')
        L.append('    assert not cur.finish(), cur.finish()   # 非空=模板被破坏，回滚')
    else:
        L.append('    先人工确认目标，再照上面四步走。')
    return '\n'.join(L)


def largest_empty_cell(frame):
    """框里最空的那个单元格，通常就是要填写的区域。返回 (row, col, 已用字数)。"""
    best = None
    for ri, tr in enumerate(frame._tbl.findall(qn('w:tr'))):
        for ci, tc in enumerate(_row_cells(tr)):
            used = len(_tc_text(tc))
            if best is None or used < best[2]:
                best = (ri, ci, used)
    return best


def survey(doc, max_rows=16):
    """勘察整份文档，给出判断写入位置所需的全部事实。

    不做判断，只摆事实：每张表的行列、边框、每行的标签与字数，并标出
    「近乎为空的格子」和「疑似评分表」。**把结论留给调用方**——模板形态太多，
    自动猜一定会在某一门上出错。
    """
    L = []
    paras = [e for e in doc.element.body.iterchildren() if e.tag == qn('w:p')]
    allt = iter_tables(doc.element.body)
    L.append('文档勘察：%d 个顶层段落、%d 张顶层表格、%d 张内嵌表格' % (
        len(paras), len(tables(doc)), len(allt) - len(tables(doc))))
    L.append('')
    for label, el in allt:
        trs = el.findall(qn('w:tr'))
        cols = _effective_cols(el)
        ncol = max(cols) if len(cols) == 1 else '/'.join(str(c) for c in sorted(cols))
        marks = []
        if looks_like_rubric(el):
            marks.append('[!] 疑似评分表 —— 不要往这里写正文')
        if cols == {1}:
            marks.append('单列长表（可被 find_frame 自动识别）')
        if '内嵌表' in label:
            marks.append('[!] 内嵌表格（题目表/信息表）—— 不要清空它所在的格子')
        L.append('%s  %d行 × %s列  %s  %s' % (
            label, len(trs), ncol,
            '带边框' if _is_bordered(el) else '无边框',
            '  '.join(marks)))
        for ri, tr in enumerate(trs[:max_rows]):
            cells = _row_cells(tr)
            desc = []
            for ci, tc in enumerate(cells):
                t = _tc_text(tc)
                desc.append('%d字%s' % (len(t), ('「%s」' % t[:22]) if t else '(空)'))
            empty = any(len(_tc_text(tc)) <= 6 for tc in cells)
            L.append('    row%-2d  %s%s' % (
                ri, ' | '.join(desc), '   <-- 近乎为空' if empty else ''))
        if len(trs) > max_rows:
            L.append('    ...（共 %d 行）' % len(trs))
        L.append('')
    if not tables(doc):
        L.append('没有顶层表格 —— 这是**段落式模板**：正文直接写在一串段落里。')
        L.append('用 Target.after("锚点文字") 指定插入位置，例如锚点为"主要内容："。')
    return '\n'.join(L)


def editable_cell(frame, row, col=0):
    """取框里第 row 行第 col 列的单元格——也就是本次要填写的位置。"""
    return frame.rows[row].cells[col]


def clear_cell(cell, keep=0):
    """清掉单元格里的段落，保留前 keep 段（模板原有的章节标题要留着）。
    只删段落，不删单元格本身，框才不会被破坏。"""
    ps = cell._tc.findall(qn('w:p'))
    for p in ps[keep:]:
        cell._tc.remove(p)


def put(cell, el):
    """把一个块级元素（或元素列表）追加进单元格。内容必须按文档顺序依次追加。

    `table_block()` / `wrapped_table()` 返回的是列表（表题段 + 表格 + 表后空段），
    这里统一展开，调用方不用自己循环。
    """
    for item in (el if isinstance(el, (list, tuple)) else [el]):
        cell._tc.append(item)


def finish_cell(cell):
    """收尾：Word 要求单元格不能以表格结尾，否则打开时报“文件内容已经损坏”。
    注意这个错误 python-docx 读不出来，只有 Word 才报。"""
    kids = [c for c in cell._tc.iterchildren() if c.tag in (qn('w:p'), qn('w:tbl'))]
    if kids and kids[-1].tag == qn('w:tbl'):
        cell._tc.append(new_p(tiny=True))


# --------------------------------------------------------------------------
# 安全写入：明确目标 + 写完校验没破坏模板
#
# "框"是什么形态，无法可靠自动识别（四种形态见 find_frame 的说明）。所以：
#   ① 目标一律【显式指定】，不猜；
#   ② 写入前 check_target 拦住高危目标（评分表、已有内容的格子）；
#   ③ 写入后 diff_snapshot 比对表格结构，确认除了目标格以外什么都没动。
# 第③步是关键——它不依赖任何猜测，直接验证结果。
# --------------------------------------------------------------------------


class Target:
    """写入目标。三种形态对应三类模板，**必须显式给出**。

        Target.cell(table=1, row=3)     表格式：往 TABLE[1] 第 3 行追加
        Target.after('主要内容：')       段落式：在含该文字的段落后插入
        Target.at_end()                 追加式：写到文档末尾
    """

    def __init__(self, mode, **kw):
        self.mode = mode
        self.kw = kw

    @classmethod
    def cell(cls, table, row, col=0):
        return cls('cell', table=table, row=row, col=col)

    @classmethod
    def after(cls, anchor, occurrence=1):
        return cls('after', anchor=anchor, occurrence=occurrence)

    @classmethod
    def at_end(cls):
        return cls('end')

    def __repr__(self):
        if self.mode == 'cell':
            return 'Target.cell(table=%d, row=%d, col=%d)' % (
                self.kw['table'], self.kw['row'], self.kw['col'])
        if self.mode == 'after':
            return 'Target.after(%r)' % self.kw['anchor']
        return 'Target.at_end()'


def find_paragraph(doc, needle, occurrence=1):
    """按文字找一个段落（返回 Paragraph），找不到返回 None。"""
    from docx.text.paragraph import Paragraph
    hits = 0
    for el in doc.element.body.iter(qn('w:p')):
        p = Paragraph(el, doc._body)
        if needle in p.text:
            hits += 1
            if hits == occurrence:
                return p
    return None


def check_target(doc, target):
    """写入前的安全检查，返回问题清单。**非空就必须停下来向用户确认。**

    拦的是四类真正的危险：往评分表里写、往封面/表头字段里写（含把正文挂在
    封面字段后面的 `Target.after()`）、清空内嵌着表格的格子、往模板已有内容里写。
    """
    problems = []
    if target.mode == 'cell':
        tbls = tables(doc)
        ti, ri, ci = target.kw['table'], target.kw['row'], target.kw['col']
        if not (0 <= ti < len(tbls)):
            return ['TABLE[%d] 不存在；文档共 %d 张顶层表格' % (ti, len(tbls))]
        el = tbls[ti]
        trs = el.findall(qn('w:tr'))
        if not (0 <= ri < len(trs)):
            return ['TABLE[%d] 没有 row%d（共 %d 行）' % (ti, ri, len(trs))]
        cells = _row_cells(trs[ri])
        if not (0 <= ci < len(cells)):
            return ['TABLE[%d] row%d 没有 col%d（该行 %d 格）' % (ti, ri, ci, len(cells))]
        if row_looks_like_rubric(el, ri):
            problems.append(
                'TABLE[%d] row%d 疑似【评分/评价行】——这一带是打分用的，'
                '往这里写正文会破坏模板。确认这真的是要填写的栏目吗？' % (ti, ri))
        rs = _rubric_start(el)
        if rs is not None and ri >= rs and not row_looks_like_rubric(el, ri):
            problems.append(
                'TABLE[%d] row%d 位于【评分区】（从 row%d 起）——'
                '评分点、赋分表、教师评语一律不许改，连空白也不行。' % (ti, ri, rs))
        if _is_cover_row(cells) or _is_cover_row([cells[ci]]):
            problems.append(
                'TABLE[%d] row%d 是【封面/表头信息行】（姓名、学号、题目…）。'
                'put() 会先清空这一格，把「姓　名：」这样的标签连同下划线一起删掉。'
                '封面字段请改用 K.fill_cover() / K.fill_field()。' % (ti, ri))
        inner = cells[ci].findall(qn('w:tbl'))
        if inner:
            problems.append(
                '目标格 TABLE[%d] row%d col%d 里**内嵌了 %d 张表格**（题目表/信息表最常见）。'
                'put() 会先清空整格：表格上方的段落会消失、表格被顶到格子最前面。'
                '改用 Target.after("表格附近的锚点文字")，或 Cursor(..., clear=False) 追加。'
                % (ti, ri, ci, len(inner)))
        used = len(_tc_text(cells[ci]))
        if used > 60:
            problems.append(
                '目标格 TABLE[%d] row%d col%d 已有 %d 字，看起来是【模板原有内容】。'
                '确认要往里写吗？' % (ti, ri, ci, used))
    elif target.mode == 'after':
        p = find_paragraph(doc, target.kw['anchor'], target.kw['occurrence'])
        if p is None:
            problems.append('找不到锚点段落：%r' % target.kw['anchor'])
            return problems
        loc = locate_cell(doc, p._p)
        if loc is not None:
            ti, ri, _ = loc
            el = tables(doc)[ti]
            rs = _rubric_start(el)
            if row_looks_like_rubric(el, ri) or (rs is not None and ri >= rs):
                problems.append(
                    '锚点段落位于 TABLE[%d] row%d 的【评分区】——'
                    '不能在评分表内部插入内容。' % (ti, ri))
        else:
            txt = _norm(p.text)
            if sum(1 for h in RUBRIC_HINTS if _norm(h) in txt) >= 2:
                problems.append(
                    '锚点段落「%s」疑似【评分/评价文字】——不要在评分区插入内容。'
                    % p.text.strip()[:20])
        # 封面/表头字段：锚点长在这些字段上时，往后插正文会把正文挤进封面。
        tc = _ancestor_tc(p._p)
        hit = _label_value(p.text)[0] or (
            _bare_label(_tc_text(tc)) if tc is not None else None)
        if hit:
            problems.append(
                '锚点段落「%s」落在【封面/表头字段】「%s」上——在它后面插正文会'
                '挤进封面。填封面字段请用 K.fill_cover()；正文换一个锚点。'
                % (target.kw['anchor'], hit))
    elif target.mode == 'end':
        blocks = [e for e in doc.element.body.iterchildren()
                  if e.tag in (qn('w:p'), qn('w:tbl'))]
        if blocks and blocks[-1].tag == qn('w:tbl') and looks_like_rubric(blocks[-1]):
            problems.append(
                '文档以【评分表】结尾，at_end() 会把正文写到评分表之后。'
                '请显式指定正文区的 Target。')
    return problems


def _table_profile(el):
    """单张表的结构指纹：行数、每行各格字数、有无边框。"""
    trs = el.findall(qn('w:tr'))
    rows = tuple(tuple(len(_tc_text(tc)) for tc in _row_cells(tr)) for tr in trs)
    return (len(trs), rows, _is_bordered(el))


def snapshot(doc):
    """表格结构指纹，**含内嵌表格**。破坏模板最典型的表现就是表格被改动，
    所以盯住表格（外框、题目表、信息表、评分表全在里面）。

    返回 `[(元素, 标签, 指纹), ...]` —— 带上元素本身，才能分辨"是不是同一张表"；
    只按序号比对会误判：往正文中间插一张表，后面所有表的序号都会错位。
    标签形如 `TABLE[1] row3 col0 内嵌表[0]`，一眼看出动的是哪一张。
    """
    return [(el, label, _table_profile(el))
            for label, el in iter_tables(doc.element.body)]


def _entry(e):
    """兼容 `(元素, 指纹)` 与 `(元素, 标签, 指纹)` 两种形态。"""
    return (e[0], None, e[1]) if len(e) == 2 else (e[0], e[1], e[2])


def diff_snapshot(before, after, texts=True):
    """比对写入前后的表格结构，返回**问题清单**（空 = 模板没被破坏）。

    只报告**原本就存在**的表格发生的变化。写入过程中新增的表格不算问题 ——
    那正是要做的操作（插入题注表、附录表等）。

    texts=False 时只比结构（行数/边框/表格增删），不比字数 ——
    专给"填封面字段"这种**预期会改字数**的操作当闸门用（见 `fill_cover`）。
    """
    before_l = [_entry(e) for e in before]
    after_l = [_entry(e) for e in after]
    before_prof = {id(el): (label, prof) for el, label, prof in before_l}
    labels = {id(el): (label or 'TABLE[#%d]' % i)
              for i, (el, label, _) in enumerate(before_l)}
    problems = []
    seen = set()
    for el, label, prof in after_l:
        key = id(el)
        if key not in before_prof:
            continue                                  # 新插入的表，正常
        seen.add(key)
        name = labels.get(key) or label or '原有表格'
        b = before_prof[key][1]
        if b[2] != prof[2]:
            problems.append('%s 的边框设置被改了' % name)
        if b[0] != prof[0]:
            problems.append('%s 行数变了：%d → %d' % (name, b[0], prof[0]))
            continue
        if texts and b[1] != prof[1]:
            for ri, (br, ar) in enumerate(zip(b[1], prof[1])):
                if br != ar:
                    problems.append('%s row%d 各格字数 %s → %s' % (
                        name, ri, list(br), list(ar)))
    gone = set(before_prof) - seen
    if gone:
        problems.append('有 %d 张原有表格不见了（被删除或替换）：%s' % (
            len(gone), '、'.join(labels.get(k, '?') for k in list(gone)[:3])))
    return problems


def _direct_cell_text(tc):
    """只取单元格自己的段落文字，不把内嵌表格里的文字算进栏目名。"""
    parts = []
    for p in tc.findall(qn('w:p')):
        parts.extend(t.text or '' for t in p.iter(qn('w:t')))
    return ''.join(parts).strip()


def _cell_shape(tc):
    """单个格子的结构：合并跨度、纵向合并、归一化后的文字（栏目名）。"""
    tcPr = tc.find(qn('w:tcPr'))
    span = 1
    vmerge = ''
    if tcPr is not None:
        gs = tcPr.find(qn('w:gridSpan'))
        if gs is not None:
            try:
                span = int(gs.get(qn('w:val'), '1'))
            except ValueError:
                span = 1
        vm = tcPr.find(qn('w:vMerge'))
        if vm is not None:
            vmerge = vm.get(qn('w:val'), 'continue')
    return (span, vmerge, _norm(_direct_cell_text(tc)))


def template_shape(doc):
    """模板结构快照：每张表的行数、每行格数、合并关系、每格栏目名。

    `snapshot()` 只比“字数”，无法发现“把一行拆成两行但字数不变”“改了左侧栏目名”
    这类结构改动。本函数专门锁死模板的**版式骨架**，写入前后必须一致。

    返回 `[(表格元素, 标签, 骨架), ...]`；骨架是
    `(行数, ((格数, ((span, vmerge, 栏目名), ...)), ...))`。
    用元素身份而不是序号比对，插入新表不会让后面的表错位。
    """
    shape = []
    for label, el in iter_tables(doc.element.body):
        rows = []
        for tr in el.findall(qn('w:tr')):
            cells = tuple(_cell_shape(tc) for tc in _row_cells(tr))
            rows.append((len(cells), cells))
        shape.append((el, label, (len(rows), tuple(rows))))
    return shape


def diff_template_shape(before, after, skip_tc=None):
    """比对模板骨架。返回问题清单（空 = 行数/列数/合并/栏目名都没变）。

    `skip_tc` 是本次唯一允许写入的目标格：只放行它自己的栏目名变化，
    它的合并属性仍然受保护。
    """
    before_map = {id(el): (label, shape) for el, label, shape in before}
    labels = {id(el): (label or 'TABLE[#%d]' % i)
              for i, (el, label, _) in enumerate(before)}
    problems = []
    seen = set()
    for el, label, shape in after:
        key = id(el)
        if key not in before_map:
            continue                                  # 新插入的表，正常
        seen.add(key)
        name = labels.get(key) or label or '原有表格'
        old_rows = before_map[key][1][0]
        new_rows = shape[0]
        if old_rows != new_rows:
            problems.append('%s 行数变了：%d → %d（模板行数必须保持原样）'
                            % (name, old_rows, new_rows))
            continue
        for ri, (old_row, new_row) in enumerate(zip(before_map[key][1][1],
                                                    shape[1])):
            if old_row[0] != new_row[0]:
                problems.append('%s row%d 列数变了：%d → %d'
                                % (name, ri, old_row[0], new_row[0]))
                continue
            for ci, (old_cell, new_cell) in enumerate(zip(old_row[1], new_row[1])):
                if old_cell[0] != new_cell[0]:
                    problems.append('%s row%d col%d 横向合并变了：%d → %d'
                                    % (name, ri, ci, old_cell[0], new_cell[0]))
                if old_cell[1] != new_cell[1]:
                    problems.append('%s row%d col%d 纵向合并变了：%r → %r'
                                    % (name, ri, ci, old_cell[1], new_cell[1]))
                if old_cell[2] != new_cell[2]:
                    tcs = _row_cells(el.findall(qn('w:tr'))[ri])
                    if skip_tc is not None and ci < len(tcs) and tcs[ci] is skip_tc:
                        continue
                    problems.append('%s row%d col%d 栏目名变了：%r → %r'
                                    % (name, ri, ci, old_cell[2], new_cell[2]))
    gone = set(before_map) - seen
    if gone:
        problems.append('有 %d 张原有表格不见了（被删除或替换）：%s' % (
            len(gone), '、'.join(labels.get(k, '?') for k in list(gone)[:3])))
    return problems


def diff_template_documents(template_doc, report_doc):
    """跨文档比对模板骨架（模板文件 vs 交付文件）。

    `diff_template_shape()` 按元素身份匹配，只适合同一份文档的写入前后；
    这里按“行数 + 列结构 + 栏目名”匹配，不会因为重新打开文件而误报。

    返回问题清单（空 = 骨架一致）。多出来的数据表只给提示，不判错。
    """
    tpl = [(label, shape) for _el, label, shape in template_shape(template_doc)]
    rep = [(label, shape) for _el, label, shape in template_shape(report_doc)]

    def pattern(shape):
        return tuple(cell_count for cell_count, _cells in shape[1])

    def labels(shape):
        return tuple(cell[2] for _count, cells in shape[1] for cell in cells)

    def score(a, b):
        return sum(1 for x, y in zip(labels(a), labels(b)) if x and x == y)

    used: set[int] = set()
    issues: list[str] = []

    def find(shape, same_rows):
        best, best_score = None, -1
        for j, (_label, cand) in enumerate(rep):
            if j in used:
                continue
            if same_rows:
                if cand[0] != shape[0] or pattern(cand) != pattern(shape):
                    continue
            else:
                if cand[0] == shape[0] or not cand[1] or not shape[1]:
                    continue
                if len(cand[1][0][1]) != len(shape[1][0][1]):
                    continue
            s = score(shape, cand)
            if any(labels(shape)) and s == 0:
                continue
            if s > best_score:
                best, best_score = j, s
        return best

    for i, (label, shape) in enumerate(tpl):
        name = label or 'TABLE[#%d]' % i
        j = find(shape, True)
        if j is None:
            j = find(shape, False)
            if j is not None:
                used.add(j)
                issues.append('%s 行数变了：%d → %d（模板行数必须保持原样）'
                              % (name, shape[0], rep[j][1][0]))
            else:
                issues.append('%s 模板表格缺失或被替换' % name)
            continue
        used.add(j)
        cand = rep[j][1]
        for ri, ((old_count, old_cells), (new_count, new_cells)) in enumerate(
                zip(shape[1], cand[1])):
            if old_count != new_count:
                issues.append('%s row%d 列数变了：%d → %d'
                              % (name, ri, old_count, new_count))
                continue
            for ci, (old_cell, new_cell) in enumerate(zip(old_cells, new_cells)):
                if old_cell[0] != new_cell[0]:
                    issues.append('%s row%d col%d 横向合并变了：%d → %d'
                                  % (name, ri, ci, old_cell[0], new_cell[0]))
                if old_cell[1] != new_cell[1]:
                    issues.append('%s row%d col%d 纵向合并变了：%r → %r'
                                  % (name, ri, ci, old_cell[1], new_cell[1]))
                if old_cell[2] and old_cell[2] != new_cell[2]:
                    issues.append('%s row%d col%d 栏目名变了：%r → %r'
                                  % (name, ri, ci, old_cell[2], new_cell[2]))
    extra = len(rep) - len(used)
    if extra > 0:
        issues.append('报告比模板多出 %d 张表（数据表/附录表正常；若接近图表数，'
                      '先查 1×1 外层表）' % extra)
    return issues


# --------------------------------------------------------------------------
# 保护区：封面、评分点、评分表、模板原有段落/表格属性、页眉页脚/样式部件
#
# diff_snapshot 只比“字数”，看不出空白、换行、字体、结构的改动。封面和评分表
# 恰恰是**连空行和空格都不能动**的地方——老师一眼就能看出排版被重排过。
# 所以这里直接对 XML 做字节级哈希：任何空白、换行、rPr/pPr、顺序变化都会被抓到。
#
# 只保护**写入前就存在**的元素；写入过程中新增的段落、表格、单元格正常忽略。
# --------------------------------------------------------------------------


def _xml_hash(el):
    """元素 XML 的哈希。空白、换行、格式属性都参与比较。"""
    from lxml import etree
    return hashlib.sha256(etree.tostring(el)).hexdigest()


def protected_snapshot(doc, skip=None):
    """记录所有**不许改动**的模板元素及其 XML 哈希。

    skip = (table_idx, row_idx, col_idx)：目标单元格。清空/写入它会改动这一格，
    所以比较该行时把这一格换成空壳，其余部分——包括格与格之间的空白——照常保护。
    返回 `[(label, element, hash), ...]`，element 用来在写入后按身份找回同一个节点。
    """
    items = []
    body = doc.element.body
    pi = ti = 0

    for child in body.iterchildren():
        tag = child.tag
        if tag == qn('w:p'):
            items.append(('P[%d] 段落' % pi, child, _xml_hash(child)))
            pi += 1
        elif tag == qn('w:tbl'):
            tbl_label = 'TABLE[%d]' % ti
            pr = child.find(qn('w:tblPr'))
            if pr is not None:
                items.append((tbl_label + ' 表格属性', pr, _xml_hash(pr)))
            grid = child.find(qn('w:tblGrid'))
            if grid is not None:
                items.append((tbl_label + ' 列定义', grid, _xml_hash(grid)))

            target_row = skip[1] if skip and skip[0] == ti else None
            target_col = skip[2] if skip and skip[0] == ti else None
            for ri, tr in enumerate(child.findall(qn('w:tr'))):
                label = '%s row%d' % (tbl_label, ri)
                if ri == target_row:
                    clone = copy.deepcopy(tr)
                    tcs = _row_cells(clone)
                    ci = target_col if target_col < len(tcs) else None
                    if ci is not None:
                        old = tcs[ci]
                        shell = OxmlElement('w:tc')
                        shell.tail = old.tail          # 留住格与格之间的空白
                        old.getparent().replace(old, shell)
                    items.append((label, tr, _xml_hash(clone)))
                else:
                    items.append((label, tr, _xml_hash(tr)))
            ti += 1
        elif tag == qn('w:sectPr'):
            items.append(('节属性（页边距/纸张）', child, _xml_hash(child)))

    if skip:
        tbls = tables(doc)
        ti_, ri_, ci_ = skip
        if ti_ < len(tbls):
            trs = tbls[ti_].findall(qn('w:tr'))
            if ri_ < len(trs):
                tcs = _row_cells(trs[ri_])
                if ci_ < len(tcs):
                    tc_pr = tcs[ci_].find(qn('w:tcPr'))
                    if tc_pr is not None:
                        items.append(('目标格的单元格属性', tc_pr, _xml_hash(tc_pr)))

    # 页眉/页脚、样式、编号等是独立的部件，不在 body 里；一并做字节级保护。
    protected_types = ('/header', '/footer', '/styles', '/numbering')
    for rel in doc.part.rels.values():
        if getattr(rel, 'is_external', False):
            continue
        if not any(rel.reltype.endswith(t) for t in protected_types):
            continue
        try:
            part = rel.target_part
            items.append(('关联部件 %s' % part.partname, part,
                          hashlib.sha256(part.blob).hexdigest()))
        except Exception:                            # noqa: BLE001
            continue
    return items


def locate_cell(doc, el):
    """找出元素所在的**顶层表格单元格**，返回 (table_idx, row_idx, col_idx)。

    元素不在表格里（或只在嵌套表里）时返回 None。`Target.after()` 的锚点
    有可能落在表格单元格内，这时插入内容会改变那一行；必须把该格排除出
    保护区，否则 finish() 会误报。
    """
    top_tbls = tables(doc)
    node = el
    while node is not None:
        tc = node.getparent()
        while tc is not None and tc.tag != qn('w:tc'):
            tc = tc.getparent()
        if tc is None:
            return None
        tr = tc.getparent()
        while tr is not None and tr.tag != qn('w:tr'):
            tr = tr.getparent()
        if tr is None:
            return None
        tbl = tr.getparent()
        while tbl is not None and tbl.tag != qn('w:tbl'):
            tbl = tbl.getparent()
        if tbl is None:
            return None

        if any(tbl is t for t in top_tbls):
            ti = next(i for i, t in enumerate(top_tbls) if tbl is t)
            ri = next((i for i, t in enumerate(tbl.findall(qn('w:tr')))
                       if t is tr), None)
            ci = next((i for i, t in enumerate(_row_cells(tr))
                       if t is tc), None)
            if ri is None or ci is None:
                return None
            return (ti, ri, ci)

        # 命中的是嵌套表：继续往上找外层单元格，保护真正会被改动的顶层行。
        node = tbl
    return None


def diff_protected(before, after):
    """比对保护区。任何空白/换行/字体/结构变化都会报出来（空 = 没动过）。"""
    after_map = {id(el): (label, h) for label, el, h in after}
    problems = []
    for label, el, h in before:
        key = id(el)
        if key not in after_map:
            problems.append('%s 不见了（被删除或替换）——模板原有内容不许删改' % label)
        elif after_map[key][1] != h:
            problems.append(
                '%s 被改动（内容、空白/换行、字体或结构都算）——'
                '封面/评分表/模板原有内容必须逐字节保持原样' % label)
    return problems


def diff_structure(before, after):
    """只比结构，不比字数 —— 填封面字段时用它验收。"""
    return diff_snapshot(before, after, texts=False)


class Cursor:
    """按显式目标往模板里落内容，并在收尾时验证模板没被破坏。

        cur = K.Cursor(doc, K.Target.cell(1, 3), keep_paras=1)
        cur.put(K.heading(1, '1. …'))      # 首次写入时自动跑安全检查
        cur.put(K.body(caps, '……'))
        cur.put_image('pics/x.png', 400)
        print(cur.finish())                # 模板结构比对报告

    **纪律是强制的，不靠自觉：**

    1. `check_target()` 在**第一次 put() 时自动执行**，拦到问题就抛错，
       写不下去。想放行必须 `cur.force('理由')` —— 理由会写进报错信息，
       逼你说清楚为什么确认这个目标是对的。
    2. 清空目标格**延迟到首次写入之后**才做。构造 Cursor 不改动文档，
       所以检查不过、你放弃写入时，模板还是原封不动的。
    3. `finish()` 比对结构指纹（含内嵌表格）**＋ 保护区 XML 哈希**，报告除
       目标格以外的任何变化——包括封面/评分表/模板原有段落里改动的一个空格
       或一个换行。

    **目标格里已经有东西时用 `clear=False`**（追加，不清空）——尤其是格子内
    嵌着题目表/信息表的时候。清空会把表格上方的段落删掉、把表格顶到最前面，
    结构指纹里的行数和字数都会对不上。
    """

    def __init__(self, doc, target, keep_paras=0, clear=True, force=None):
        self.doc = doc
        self.target = target
        self.parent = doc._body
        self.before = snapshot(doc)
        # 版式骨架：行数/列数/合并/栏目名。比 snapshot 更严格，专门防“加行/改栏目名”。
        self.shape_before = template_shape(doc)
        self._anchor_el = None
        self._keep_paras = keep_paras
        self._clear = clear
        self._cleared = False
        self._cell_tc = None
        self._write_tc = None                 # 允许写入的单元格（含 Target.after 所在格）
        self._problems = None                 # 惰性求值，首次 put 时算
        self._reason = force
        self._skip = None

        if target.mode == 'cell':
            tbls = tables(doc)
            el = tbls[target.kw['table']]
            tr = el.findall(qn('w:tr'))[target.kw['row']]
            self._cell_tc = _row_cells(tr)[target.kw['col']]
            self._write_tc = self._cell_tc
            self._skip = (target.kw['table'], target.kw['row'], target.kw['col'])
            # 注意：这里**不清空**。清空推迟到首次 put()，见 _guard()。
        elif target.mode == 'after':
            p = find_paragraph(doc, target.kw['anchor'], target.kw['occurrence'])
            if p is None:
                raise ValueError('找不到锚点段落：%r' % target.kw['anchor'])
            self._anchor_el = p._p
            # 锚点在表格单元格内时，插入会改变该格；把它排除出保护区。
            self._skip = locate_cell(doc, p._p)
            self._write_tc = _ancestor_tc(p._p)
        else:                                     # end：用 body 里最后一个块级元素作锚
            blocks = [e for e in doc.element.body.iterchildren()
                      if e.tag in (qn('w:p'), qn('w:tbl'))]
            self._anchor_el = blocks[-1] if blocks else None

        # 保护区必须在任何改动之前拍快照；目标格本身允许变，其余逐字节保护。
        self._protected = protected_snapshot(doc, self._skip)

    # ---------------------------------------------------------------- 安全检查
    def check(self):
        """目标安全检查，返回问题清单（空 = 没问题）。首次 put() 会自动调它。"""
        if self._problems is None:
            self._problems = check_target(self.doc, self.target)
        return self._problems

    def force(self, reason):
        """确认目标无误、显式放行。**必须给理由**，它会进日志，也会跟着报错一起显示。

        只在 `check()` 报了问题、而你确实判断那是个误报时才用。
        """
        if not str(reason or '').strip():
            raise ValueError('force() 必须给一个理由：写清为什么确认这个目标是对的')
        self._reason = str(reason)
        return self.check()

    def resnapshot(self, reason):
        """用户**明确同意**改模板结构后，重新拍快照，作为后续比对的基准。

        这不是绕过检查：默认仍然禁止加行/改栏目名；只有用户看过影响并明确同意，
        才调用本方法，并把理由写下来（理由会打印出来）。
        """
        if not str(reason or '').strip():
            raise ValueError('resnapshot() 必须给理由：用户同意改哪一处结构、为什么')
        self.before = snapshot(self.doc)
        self.shape_before = template_shape(self.doc)
        self._protected = protected_snapshot(self.doc, self._skip)
        print('[docx_kit] 已按用户许可重新快照模板结构：%s' % reason)
        return self

    def _guard(self):
        """首次写入前的闸门。检查没过又没有理由 → 直接抛错，什么都不改。"""
        problems = self.check()
        if problems and not self._reason:
            raise ValueError(
                '写入被拦截 —— 目标看起来不对：\n    - %s\n'
                '\n文档未被改动。请换一个 Target；若确认无误，'
                '调用 cur.force("理由") 显式放行。' % '\n    - '.join(problems))
        if problems:
            print('[docx_kit] 已按理由放行：%s\n           目标：%r'
                  % (self._reason, self.target))
        return problems

    # ---------------------------------------------------------------- 写入
    def put(self, el):
        """落一个块级元素（或元素列表）。首次调用会自动跑安全检查。

        `table_block()` 返回列表（表题段 + 表格 + 表后空段），可直接 put。
        """
        if isinstance(el, (list, tuple)):
            for item in el:
                self.put(item)
            return el[-1] if el else None
        self._guard()
        if self.target.mode == 'cell':
            if not self._cleared:                 # 通过了检查才清空目标格
                if self._clear:
                    clear_cell_tc(self._cell_tc, keep=self._keep_paras)
                self._cleared = True
            self._cell_tc.append(el)
        elif self._anchor_el is None:                 # 空文档：插到 sectPr 之前
            self.doc.element.body.insert(0, el)
            self._anchor_el = el
        else:
            self._anchor_el.addnext(el)
            self._anchor_el = el
        return el

    def put_image(self, path, width_pt):
        return self.put(image(self.parent, path, width_pt))

    def put_table(self, caps, key, cap_text, rows, widths, **kw):
        """一次落“表题 + 表格 + 表后空行”，不含 1×1 外层表。"""
        return self.put(table_block(caps, key, cap_text, rows, widths, **kw))

    # ---------------------------------------------------------------- 收尾
    def finish(self):
        """收尾并返回模板保护报告（列表。空 = 没动到模板其它地方）。

        三道检查一起报：表格结构指纹（含内嵌表格，字数级）+ 模板骨架
        （行数/列数/合并/栏目名）+ 保护区 XML 哈希（字节级，能抓到空白、
        换行、字体和结构改动）。
        """
        if self._cell_tc is not None:
            finish_cell(_CellProxy(self._cell_tc))
        elif self._anchor_el is not None:
            # 锚点在某个单元格里（往内嵌表格附近的段落后面插）时，同一个规则
            # 也要守：格子里最后一块不能是表格，否则 Word 报"文件内容已经损坏"。
            tc = _ancestor_tc(self._anchor_el)
            if tc is not None:
                finish_cell(_CellProxy(tc))
        problems = diff_snapshot(self.before, snapshot(self.doc))
        skip_tc = self._cell_tc if self._cell_tc is not None else self._write_tc
        problems += diff_template_shape(self.shape_before, template_shape(self.doc),
                                        skip_tc=skip_tc)
        problems += diff_protected(self._protected, protected_snapshot(self.doc, self._skip))
        return problems


class _CellProxy:
    """给 finish_cell 用的最小壳子（只要 ._tc）。"""

    def __init__(self, tc):
        self._tc = tc


def _ancestor_tc(el):
    """往上找到最近的 `w:tc`（元素不在表格里就返回 None）。"""
    p = el.getparent()
    while p is not None:
        if p.tag == qn('w:tc'):
            return p
        p = p.getparent()
    return None


def clear_cell_tc(tc, keep=0):
    """按 tc 元素清段落（clear_cell 的底层版本）。

    ⚠️ 只删段落，**不删表格**——所以格子里若原本有内嵌表格（题目表/信息表），
    清空之后它会被顶到格子最前面，上面原本的段落全没了。别对含内嵌表格的格子
    调这个（`check_target()` 会拦）；改用 `Target.after("锚点")` 插到表格后面。
    """
    ps = tc.findall(qn('w:p'))
    for p in ps[keep:]:
        tc.remove(p)


# --------------------------------------------------------------------------
# 封面 / 表头字段：只填值，不动标签
#
# 封面——以及题目表、学生信息表里那些「姓　名：____」「题　目」——**不是正文编写区**。
# 拿 Cursor/Target.cell 去写它，等于先清空整格：标签、下划线、校徽、表格边框一起没。
# 所以另起一套：
#   ① cover_fields()  只勘察：哪个字段、现在是什么值、值该落在哪
#   ② 缺哪些值 **问用户**（或让用户自己手填），姓名/学号/日期绝不猜
#   ③ fill_cover()/fill_field()  填值——复用模板自己留的空白 run，其余一字不动
# --------------------------------------------------------------------------

_BLANKS = '_＿ \t　'


def _as_p(p):
    """段落：既能收 `docx.text.paragraph.Paragraph`，也能收 `<w:p>` 元素。"""
    return getattr(p, '_p', p)


def _p_text(p):
    return ''.join(t.text or '' for t in _as_p(p).iter(qn('w:t')))


def _tbl_first_text(tbl):
    trs = tbl.findall(qn('w:tr'))
    return _tc_text(_row_cells(trs[0])[0]) if trs else ''


def _label_value(text):
    """把「题　目：____」拆成 `('题目', '')`；不是「标签：值」就返回 `(None, None)`。

    标签要先去掉全角空格再认，因为模板爱写成「姓　名」「班　　级」。
    值里的下划线/空格也一并去掉——那只是给手写留的空白，不是已填的值。
    """
    s = (text or '').strip()
    m = re.match(r'^([^：:\n]{1,12})[：:][\s　]*([\s\S]*)$', s)
    if not m:
        return None, None
    label = _norm(m.group(1))
    if not label or not any(h in label for h in COVER_HINTS):
        return None, None
    return label, m.group(2).strip(_BLANKS)


def _cell_value(tc):
    """右邻格里的值。**只有下划线/空格的算"还没填"**——模板常写成
    `| 题　目： | ＿＿＿＿＿ |`，那串下划线是留给手写的空白，不是已填的内容。"""
    t = _tc_text(tc)
    return '' if t.strip(_BLANKS) == '' else t


def _bare_label(text):
    """整格就是一个标签、**没有冒号**的写法（题目表最常见：`| 题　目 | 值 |`）。"""
    s = _norm(text)
    if 0 < len(s) <= 8 and any(h in s for h in COVER_HINTS):
        return s
    return None


def _run_text(r):
    return ''.join(t.text or '' for t in r.findall(qn('w:t')))


def _is_blank_run(r):
    """这个 run 是不是"留给填的空白"：没文字、纯空格、或一串下划线。"""
    return _run_text(r).strip(_BLANKS) == ''


def _set_run_text(r, text):
    """只改这个 run 的文字，**格式一律不动**（下划线、字体全靠它保留）。"""
    ts = r.findall(qn('w:t'))
    for extra in ts[1:]:
        r.remove(extra)
    if ts:
        t = ts[0]
    else:
        t = OxmlElement('w:t')
        r.append(t)
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    return r


_AFTER_U = ('w:effect', 'w:bdr', 'w:shd', 'w:fitText', 'w:vertAlign', 'w:rtl',
            'w:cs', 'w:em', 'w:lang', 'w:eastAsianLayout', 'w:specVanish',
            'w:oMath')


def _ensure_underline(r):
    """给这个 run 加下划线。`w:u` 的位置有讲究（在 sz 之后、vertAlign 之前），
    插错位置 Word 会报"文件内容已经损坏"。"""
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        r.insert(0, rPr)
    u = rPr.find(qn('w:u'))
    if u is None:
        u = E('w:u', val='single')
        for child in rPr:
            if child.tag in [qn(x) for x in _AFTER_U]:
                child.addprevious(u)
                break
        else:
            rPr.append(u)
    else:
        u.set(qn('w:val'), 'single')
    return r


def _strip_underline(r):
    rPr = r.find(qn('w:rPr'))
    if rPr is not None:
        for u in rPr.findall(qn('w:u')):
            rPr.remove(u)
    return r


def _split_run_at(r, cut):
    """把 run 的文字从第 cut 个字符切开，后半段独立成一个 run（格式照抄）。
    用来把「姓　名：示例姓名」里的值部分单独拎出来。"""
    t = _run_text(r)
    _set_run_text(r, t[:cut])
    new = OxmlElement('w:r')
    src = r.find(qn('w:rPr'))
    if src is not None:
        new.append(copy.deepcopy(src))
    _set_run_text(new, t[cut:])
    r.addnext(new)
    return new


def _label_end(p, label):
    """段落文字里「标签：」结束的位置（字符下标）。找不到返回 None。"""
    m = re.match(r'^[\s　]*([^：:\n]{1,12})[：:]', _p_text(p))
    if not m:
        return None
    if label and _norm(label) not in _norm(m.group(1)) \
            and _norm(m.group(1)) not in _norm(label):
        return None
    return m.end()


def _tail_runs(p, label=None):
    """标签之后的那些 run —— 值该落的地方。给了 label 就只取标签后面的，
    这样**绝不会误伤标签本身**（「姓」「　」「名」「：」各占一个 run 也不怕）。"""
    p = _as_p(p)
    runs = p.findall(qn('w:r'))
    if not label:
        return runs
    cut = _label_end(p, label)
    if cut is None:
        return []
    pos = 0
    for i, r in enumerate(runs):
        n = len(_run_text(r))
        if pos + n > cut:
            return [_split_run_at(r, cut - pos)] + runs[i + 1:]
        pos += n
    return []


def _value_prop(src):
    """把标签 run 的格式给值用：字体字号照抄，**去掉加粗**（标签常加粗，值一般不加）。"""
    if src is None:
        return None
    prop = copy.deepcopy(src)
    for tag in ('w:b', 'w:bCs', 'w:i', 'w:iCs'):
        for el in prop.findall(qn(tag)):
            prop.remove(el)
    return prop


def _label_prop(para):
    runs = _as_p(para).findall(qn('w:r'))
    return _value_prop(runs[-1].find(qn('w:rPr'))) if runs else None


def _field(label, value, where, kind, para, prop):
    return {'label': label, 'value': value, 'where': where,
            'kind': kind, 'para': para, 'prop': prop}


def cover_fields(doc, limit=40):
    """勘察封面/表头里的字段，**只报告、不改动**。

    返回 `[{'label', 'value', 'where', 'kind', 'para', 'prop'}, ...]`：

    | kind | 什么情况 | 值落在哪 |
    | --- | --- | --- |
    | `inline` | 标签和值在同一段（`姓　名：____`） | 标签段本身 |
    | `next_para` | 标签独占一段，值留在下一段 | 下一个空白段落 |
    | `next_cell` | 标签独占一格，值是右邻格（题目表常见） | 右邻格首个段落 |
    | `manual` | 找不到可靠的落点 | 不动，交人工处理 |

    `value` 为空 = 这个字段还空着。**空着的要问用户**，别猜姓名/学号/日期。
    填值用 `fill_cover()`，不要用 `put()`。
    """
    out = []
    for tlabel, tbl in iter_tables(doc.element.body):
        for ri, tr in enumerate(tbl.findall(qn('w:tr'))):
            cells = _row_cells(tr)
            for ci, tc in enumerate(cells):
                ps = tc.findall(qn('w:p'))
                if not ps:
                    continue
                where = '%s row%d col%d' % (tlabel, ri, ci)
                lab = lab_para = None
                for q in ps:
                    l, v = _label_value(_p_text(q))
                    if l:
                        lab, lab_para = l, q
                        break
                if lab is None:
                    # 「题目」这种没有冒号、值在右邻格的写法
                    bare = _bare_label(_tc_text(tc))
                    if bare is None:
                        continue
                    nxt = cells[ci + 1] if ci + 1 < len(cells) else None
                    nps = nxt.findall(qn('w:p')) if nxt is not None else []
                    if not nps:
                        out.append(_field(bare, _cell_value(nxt) if nxt is not None
                                          else '', where, 'manual', None, None))
                        continue
                    out.append(_field(bare, _cell_value(nxt), where + ' →下一格',
                                      'next_cell', nps[0], _label_prop(ps[0])))
                    continue

                val = _label_value(_p_text(lab_para))[1]
                tail = _tail_runs(lab_para, lab)
                has_slot = any(_is_blank_run(r) for r in tail)
                nxt = cells[ci + 1] if ci + 1 < len(cells) else None
                nps = nxt.findall(qn('w:p')) if nxt is not None else []
                blank_para = next((q for q in ps[ps.index(lab_para) + 1:]
                                   if _p_text(q).strip() == ''), None)

                if val:
                    out.append(_field(lab, val, where, 'inline', lab_para, None))
                elif has_slot:
                    out.append(_field(lab, '', where + '（下划线处）', 'inline',
                                      lab_para, None))
                elif nxt is not None and nps:
                    # 值在右邻格：可能是已填的，也可能只留了一串下划线
                    out.append(_field(lab, _cell_value(nxt), where + ' →下一格',
                                      'next_cell', nps[0], _label_prop(lab_para)))
                elif blank_para is not None:
                    out.append(_field(lab, '', where + '（下一段）', 'next_para',
                                      blank_para, _label_prop(lab_para)))
                else:
                    out.append(_field(lab, '', where + '（标签后）', 'inline',
                                      lab_para, _label_prop(lab_para)))

    # 段落式封面：正文顶层、没有表格兜着的那种
    for pi, el in enumerate(doc.element.body.findall(qn('w:p'))):
        lab, val = _label_value(_p_text(el))
        if lab:
            out.append(_field(lab, val, '段落%d' % pi, 'inline', el,
                              _label_prop(el)))
    return out[:limit]


def fill_field(p, text, label=None, underline=None, prop=None):
    """往「标签：____」这个位置填值。**不清空、不删标签、不动版式。**

    优先级：① 标签后本来就有空白填充（下划线/空格）→ 只改它的文字，格式原样保留；
    ② 标签后已有内容 → 替换那一部分，标签一字不动；③ 标签后什么都没有 → 追加一个
    run，字体字号照抄标签（去掉加粗）。

    `underline=True` 才强制打下划线；`None`（默认）= 模板什么样就什么样。
    `p` 收 `Paragraph` 也行，收 `<w:p>` 元素也行。
    """
    p = _as_p(p)
    tail = _tail_runs(p, label)
    if tail:
        slot = tail[0]
        _set_run_text(slot, text)
        for r in tail[1:]:
            _set_run_text(r, '')
    else:
        if prop is None:
            runs = p.findall(qn('w:r'))
            if runs:
                prop = _value_prop(runs[-1].find(qn('w:rPr')))
        slot = run(text, prop)
        p.append(slot)
    if underline:
        _ensure_underline(slot)
    elif underline is False:
        _strip_underline(slot)
    return slot


def fill_cover(doc, values, underline=None, limit=40):
    """按标签把值填进封面/表头字段。**只改值，不动标签、不清空、不删框。**

        before = K.snapshot(doc)
        filled, missed, manual = K.fill_cover(doc, {'姓名': '示例姓名'})
        print(K.diff_structure(before, K.snapshot(doc)))   # 必须为空

    同一标签出现多次就都填（封面一份、题目表一份很常见）。
    返回 `(已填, 没找到的标签, 需人工处理的字段)`。
    """
    fields = cover_fields(doc, limit=limit)
    filled, missed, manual = [], [], []
    for key, val in values.items():
        k = _norm(str(key))
        if not k:
            continue
        hits = [f for f in fields
                if f['label'] == k or k in f['label'] or f['label'] in k]
        if not hits:
            missed.append(key)
            continue
        for f in hits:
            if f['kind'] == 'manual' or f['para'] is None:
                manual.append('%s（%s）' % (f['label'], f['where']))
                continue
            fill_field(f['para'], str(val),
                       label=f['label'] if f['kind'] == 'inline' else None,
                       underline=underline, prop=f['prop'])
            filled.append('%s：%s  [%s %s]' % (f['label'], val, f['where'],
                                               f['kind']))
    return filled, missed, manual


# --------------------------------------------------------------------------
# 题注编号与交叉引用（两遍法）
# --------------------------------------------------------------------------


class Captions:
    """先按最终出现顺序 declare 一遍，再开始渲染。

    为什么要两遍：正文里“如图 4 所示”可能出现在图 4 之前，一边写一边编号
    会拿到错的号。声明和渲染分开就没有这个问题。
    """

    def __init__(self):
        self._seq = {}
        self._n = {}

    def declare(self, key, kind):
        """key 自取（如 'fig_topo'），kind 是 '图' / '表' / 'Figure' 等前缀。"""
        self._seq[kind] = self._seq.get(kind, 0) + 1
        self._n[key] = '%s %d' % (kind, self._seq[kind])
        return self._n[key]

    def ref(self, key):
        if key not in self._n:
            raise KeyError('题注 key 未登记：%s' % key)
        return self._n[key]

    def number(self, key):
        return int(self._n[key].split()[-1])


_BM = [10000]


def caption(caps, key, text, style='Caption', keep_next=False):
    """题注段落：`图 3 交换机的端口映射`。

    编号用 SEQ 域（自动编号，插删图表后仍连续），外面套书签供正文 REF 引用。
    题注格式统一为“前缀 + 空格 + 编号 + 空格 + 题注名”，题注名中英文打头都一样，
    全文保持一致。图题在图下方，表题在表上方。
    """
    kind = caps.ref(key).rsplit(' ', 1)[0]
    n = caps.number(key)
    p = new_p(style=style, keep_lines=True, center=True, keep_next=keep_next)
    bid = str(_BM[0])
    _BM[0] += 1
    p.append(E('w:bookmarkStart', id=bid, name='_Ref_' + key))
    p.append(run(kind))
    p.append(run(' ', preserve=True))
    for r in field(' SEQ %s \\* ARABIC ' % kind, str(n)):
        p.append(r)
    p.append(E('w:bookmarkEnd', id=bid))
    if text:
        p.append(run(' ' + text, preserve=True))
    return p


_REF_MARK = re.compile(r'<<([A-Za-z0-9_:,\-]+)>>')


def cite(*nums):
    """生成上标形式的文献引用标注，如 ``[1]`` ``[1-3]`` ``[1,3]``。

    规范要求正文里的引用序号是**右上角标**，不是普通方括号——这点很容易漏。
    """
    parts = []
    for n in nums:
        parts.extend(str(n).split(',') if isinstance(n, str) else [str(n)])
    if len(parts) == 1:
        label = parts[0]
    else:
        ints = sorted(int(p) for p in parts if p.strip().isdigit())
        label = ('%d-%d' % (ints[0], ints[-1])) if len(ints) == len(parts) and \
            ints == list(range(ints[0], ints[-1] + 1)) else ','.join(parts)
    return run('[' + label + ']', rpr(superscript=True))


def body(caps, text, indent=True, prop=None):
    """正文段落。两种标记：

    ``<<fig_topo>>``     → 图表交叉引用（REF 域，编号自动跟随）
    ``<<cite:1>>``       → 上标文献引用 ``[1]``；``<<cite:1-3>>`` / ``<<cite:1,3>>`` 同理

    例：``K.body(caps, '该结论已被多项研究证实<<cite:1-3>>。详见<<fig_result>>。')``
    """
    p = new_p(indent=indent)
    pos = 0
    for m in _REF_MARK.finditer(text):
        if m.start() > pos:
            p.append(run(text[pos:m.start()], prop))
        tok = m.group(1)
        if tok.startswith('cite:'):
            p.append(cite(tok[5:]))
        else:
            for r in field(' REF _Ref_%s \\h ' % tok, caps.ref(tok)):
                p.append(r)
        pos = m.end()
    if pos < len(text):
        p.append(run(text[pos:], prop))
    return p


# --------------------------------------------------------------------------
# 段落块
# --------------------------------------------------------------------------


def heading(level, text, prop=None):
    """小标题。层级顺序：1. 2. 3. → (1) (2) (3) → ①②③ → 1.1 1.2 → 2.1 2.2。
    格式随模板：模板把小标题做成什么样（加粗 / 黑体 / 放大），照着做，
    不要自作主张换字体或字号。"""
    p = new_p(keep_next=True)
    p.append(run(text, prop if prop is not None else rpr(bold=True)))
    return p


def code_block(text, sz=18, font='Consolas', shade='F2F2F2', keep_next=False):
    """等宽字体 + 灰底的代码块。

    报告里贴代码的取舍：
    - 用户给了代码截图 → 直接用截图（有语法高亮，版面也省）；
    - 没给截图 → 用本函数生成文本块，正文可复制、可检索，查重时也好读。

    字号默认 sz=18（9pt，比正文五号(10.5pt)小一号）。整块是一个段落，
    行间用 <w:br/> 分隔，这样灰底连成一片，不会出现行间白缝。
    """
    # 不设段前段后：灰底本身就有区分度，而插入段落带显式间距会被一眼看出来
    p = new_p(shade=shade, keep_next=keep_next)
    lines = str(text).rstrip('\n').split('\n')
    for i, line in enumerate(lines):
        if i:
            p.append(OxmlElement('w:br'))
        p.append(run(line, rpr(fonts=font, sz=sz, hint=None), preserve=True))
    return p


def image(parent, path, width_pt, crop=None):
    """居中插图。默认带 keepNext，保证图不会被甩到页尾而题注在下一页。"""
    p = new_p(keep_next=True, center=True)
    r = OxmlElement('w:r')
    p.append(r)
    Paragraph(p, parent).runs[0].add_picture(path, width=Pt(width_pt))
    return p


# 单元格文字：9.5pt 是常见报告模板的表格字号，按模板调整
def _cell_text(txt, sz=19, bold=False, keep_next=True):
    ps = []
    for line in str(txt).split('\n'):
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        if keep_next:
            pPr.append(OxmlElement('w:keepNext'))
        pPr.append(E('w:jc', val='center'))
        p.append(pPr)
        p.append(run(line, rpr(bold=bold, sz=sz)))
        ps.append(p)
    return ps


def _tbl_pr(width, borders):
    """表格属性。

    borders:
      True        全框线（普通数据表）
      False       无边框（只给内部容器用；**不要**再用 1×1 外层表包表题）
      'three-line' 三线表：只有顶线、表头下线、底线，左右和内部竖线为 nil
    """
    pr = OxmlElement('w:tblPr')
    pr.append(E('w:tblW', w=str(width), type='dxa'))
    pr.append(E('w:jc', val='center'))
    if borders:
        b = OxmlElement('w:tblBorders')
        if borders == 'three-line':
            # 顺序必须符合 schema：top → left → bottom → right → insideH → insideV
            for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                if side in ('top', 'bottom'):
                    b.append(E('w:' + side, val='single', sz='6', space='0',
                               color='000000'))
                else:
                    b.append(E('w:' + side, val='nil'))
        else:
            for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                b.append(E('w:' + side, val='single', sz='4', space='0',
                           color='000000'))
        pr.append(b)
    pr.append(E('w:tblLayout', type='fixed'))
    # 单元格左右留白：不留白时相邻列文字会贴在一起（WPS 尤其明显）。
    ml = OxmlElement('w:tblCellMar')
    ml.append(E('w:left', w='80', type='dxa'))
    ml.append(E('w:right', w='80', type='dxa'))
    pr.append(ml)
    pr.append(E('w:tblLook', val='0000', firstRow='0', lastRow='0',
                firstColumn='0', lastColumn='0', noHBand='0', noVBand='0'))
    return pr


def _set_grid(tbl, widths):
    old = tbl.find(qn('w:tblGrid'))
    if old is not None:
        tbl.remove(old)
    g = OxmlElement('w:tblGrid')
    for w in widths:
        g.append(E('w:gridCol', w=str(w)))
    tbl.find(qn('w:tblPr')).addnext(g)


def make_table(rows, widths, sz=19, split=False, three_line=False,
               keep_together=True):
    """数据表。rows[0] 当表头。

    - `split=True`：允许整表跨页续排（单行内部仍不可拆），解决“页尾留大片空白”。
    - `three_line=True`：三线表（顶线/表头下线/底线，无左右和内部竖线）。
    - `keep_together=True`：给除最后一行外的所有行加 keepNext，让表格尽量不被拆开。
      `split=True` 时自动关闭。

    ⚠️ 本函数只返回表格本身；表题、表后空行由 `table_block()` 一并生成。
    """
    if split:
        keep_together = False
    tbl = OxmlElement('w:tbl')
    tbl.append(_tbl_pr(sum(widths), 'three-line' if three_line else True))
    tbl.append(OxmlElement('w:tblGrid'))
    n_rows = len(rows)
    for ri, row in enumerate(rows):
        tr = OxmlElement('w:tr')
        trPr = OxmlElement('w:trPr')
        trPr.append(OxmlElement('w:cantSplit'))
        if split and ri == 0:
            trPr.append(OxmlElement('w:tblHeader'))   # 嵌套在单元格里时会被 Word 忽略
        trPr.append(E('w:jc', val='center'))
        tr.append(trPr)
        for ci, txt in enumerate(row):
            tc = OxmlElement('w:tc')
            tcPr = OxmlElement('w:tcPr')
            tcPr.append(E('w:tcW', w=str(widths[ci]), type='dxa'))
            if three_line and ri == 0:
                # 三线表的中间那条线：表头行底部边框
                tb = OxmlElement('w:tcBorders')
                tb.append(E('w:bottom', val='single', sz='6', space='0',
                            color='000000'))
                tcPr.append(tb)
            tc.append(tcPr)
            for p in _cell_text(txt, sz=sz, bold=(ri == 0),
                                keep_next=(keep_together and ri < n_rows - 1)):
                tc.append(p)
            tr.append(tc)
        tbl.append(tr)
    _set_grid(tbl, widths)
    return tbl


def table_block(caps, key, cap_text, rows, widths, sz=19, split=False,
                three_line=True, blank_after=True):
    """表题段（keepNext）+ 表格 + 表后空段，返回**元素列表**，可直接 `cur.put()`。

    为什么不用 1×1 外层表：Word 里看不出来，但 WPS 会在数据表外再显示一圈
    表格边框。用户看到的就是“表格外面又套了一层表”。所以这里彻底去掉外层表。

    同页靠三件事保证：
      1. 表题段落显式 `keepNext`；
      2. 表格除最后一行外的行都带 `keepNext`（`split=True` 时关闭）；
      3. 不拆行（每行 `cantSplit`）。

    `blank_after=True` 在表格后补一个**正常空段落**：既满足“表格后空一行”，
    也满足“单元格不能以表格结尾”。不要把它换成 1pt 的 tiny 段落——
    那样只是占位，肉眼看不到“空一行”。
    """
    block = [caption(caps, key, cap_text, keep_next=True),
             make_table(rows, widths, sz=sz, split=split,
                        three_line=three_line, keep_together=not split)]
    if blank_after:
        block.append(new_p())
    return block


def wrapped_table(caps, key, cap_text, rows, widths, sz=19, split=False,
                  outer_w=None):
    """兼容旧调用的别名。**不再生成 1×1 外层表**，返回同 `table_block()` 的列表。

    `outer_w` 保留只为向后兼容，已无作用。新代码请直接写 `K.table_block(...)`。
    """
    return table_block(caps, key, cap_text, rows, widths, sz=sz, split=split)


# --------------------------------------------------------------------------
# 排版细节
# --------------------------------------------------------------------------

_CJK = ('[\\u3000-\\u303f\\u4e00-\\u9fff\\uff01-\\uff60'
        '\\u2018\\u2019\\u201c\\u201d\\u2014\\u2026]')
_SP_A = re.compile(r'(?<=%s)[ \t]+(?=[0-9A-Za-z])' % _CJK)
_SP_B = re.compile(r'(?<=[0-9A-Za-z%%])[ \t]+(?=%s)' % _CJK)


def despace(s):
    """去掉中文与英文/数字之间手打的空格，写作 `使用Console线`。

    中文档正文多为两端对齐，手打的中西文空格是唯一可被拉伸的空白，
    被拉开后行内会出现明显空洞；交给 Word 的“自动调整中文与西文间距”
    自己补视觉间隙即可。仅正文用，题注、表格（居中、短）不必处理。
    """
    return _SP_B.sub('', _SP_A.sub('', s))


def set_compatibility_mode(doc, mode=15):
    """把 `word/settings.xml` 的 compatibilityMode 设为 15（Word 2013+）。

    `.doc` 转 `.docx` 后常常仍是 11（Word 2003）。保持 11 时，Word 保存会把
    OMML 公式转成图片，而且 Word 关闭时会提示“以前版本/公式转图片”。
    这一步要在 Word 保存之后做，否则 Word 可能写回 11。
    """
    settings = doc.settings.element
    compat = settings.find(qn('w:compat'))
    if compat is None:
        compat = OxmlElement('w:compat')
        settings.append(compat)
    # 现代 Word 把模式存在 <w:compatSetting w:name="compatibilityMode">，
    # 旧写法是 <w:compatibilityMode>。两个都设置，Word 才不会忽略。
    setting = None
    for node in compat.findall(qn('w:compatSetting')):
        if node.get(qn('w:name')) == 'compatibilityMode':
            setting = node
            break
    if setting is None:
        setting = OxmlElement('w:compatSetting')
        setting.set(qn('w:name'), 'compatibilityMode')
        setting.set(qn('w:uri'), 'http://schemas.microsoft.com/office/word')
        compat.append(setting)
    setting.set(qn('w:val'), str(mode))
    node = compat.find(qn('w:compatibilityMode'))
    if node is not None:          # 旧写法存在才同步，不要凭空新增
        node.set(qn('w:val'), str(mode))
    return mode


def compatibility_mode(doc):
    """读当前 compatibilityMode；没有该节点时返回 None。"""
    settings = doc.settings.element
    compat = settings.find(qn('w:compat'))
    if compat is None:
        return None
    for node in compat.findall(qn('w:compatSetting')):
        if node.get(qn('w:name')) == 'compatibilityMode':
            try:
                return int(node.get(qn('w:val'), '0'))
            except ValueError:
                return None
    node = compat.find(qn('w:compatibilityMode'))
    if node is None:
        return None
    try:
        return int(node.get(qn('w:val'), '0'))
    except ValueError:
        return None


_M_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
_PPR_ORDER = [
    'pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr',
    'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
    'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct',
    'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
    'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents',
    'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
    'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr',
    'pPrChange',
]


def _ppr_set(pPr, tag, attrs=None):
    """在 w:pPr 里按 schema 顺序设置一个子元素（已存在则改属性）。"""
    name = tag.split(':')[-1]
    node = pPr.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        rank = _PPR_ORDER.index(name) if name in _PPR_ORDER else len(_PPR_ORDER)
        inserted = False
        for child in list(pPr):
            cname = child.tag.split('}')[-1]
            crank = _PPR_ORDER.index(cname) if cname in _PPR_ORDER else len(_PPR_ORDER)
            if crank > rank:
                child.addprevious(node)
                inserted = True
                break
        if not inserted:
            pPr.append(node)
    for key, value in (attrs or {}).items():
        node.set(qn('w:' + key), str(value))
    return node


def left_align_math_paragraphs(doc):
    """让所有含 OMML 的段落显式左对齐，并关闭 autoSpaceDE/autoSpaceDN。

    Word/WPS 会把继承 `Normal` 的两端对齐套到公式上，导致公式被拉伸、
    角标间距变大。返回处理的段落数。
    """
    changed = 0
    for p in doc.element.body.iter(qn('w:p')):
        if not (list(p.iter(_M_NS + 'oMath')) or
                list(p.iter(_M_NS + 'oMathPara'))):
            continue
        pPr = p.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p.insert(0, pPr)
        _ppr_set(pPr, 'w:jc', {'val': 'left'})
        _ppr_set(pPr, 'w:autoSpaceDE', {'val': '0'})
        _ppr_set(pPr, 'w:autoSpaceDN', {'val': '0'})
        changed += 1
    return changed


def force_western_font(doc, font='Times New Roman'):
    """给所有“纯非中文”run 写入显式西文字体（ascii/hAnsi/eastAsia/cs + hint=default）。

    WPS 对没有显式字体的希腊字母/下标会回退到宋体，按中文字宽排，表现为角标间距大。
    **必须在 Word 保存之后调用**：Word 保存会把 `w:rFonts` 改回 `hint=eastAsia`。
    含中文的 run 跳过，避免把正文中文改成西文字体。

    返回改动的 run 数量。
    """
    import re as _re
    cjk = _re.compile(r'[　-〿一-鿿！-｠]')
    western = _re.compile(r'[A-Za-zͰ-Ͽ]')
    changed = 0
    for r in doc.element.body.iter(qn('w:r')):
        text = ''.join(t.text or '' for t in r.iter(qn('w:t')))
        if not text or not western.search(text) or cjk.search(text):
            continue
        rPr = r.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            r.insert(0, rPr)
        fonts = rPr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rPr.insert(0, fonts)
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + attr), font)
        fonts.set(qn('w:hint'), 'default')
        changed += 1
    return changed


def file_in_use(path):
    """目标文件是否被 Word/WPS 占用（存在 `~$` 锁文件或无法独占打开）。

    写入前先查，别等覆盖失败才报错。返回 True 时不要反复重试，
    改用 `<名称>-修正版.<后缀>` 并在交付说明里写清楚。
    """
    p = os.path.abspath(path)
    folder, name = os.path.split(p)
    if os.path.exists(os.path.join(folder, '~$' + name)):
        return True
    if not os.path.exists(p):
        return False
    try:
        f = open(p, 'r+b')
        f.close()
        return False
    except OSError:
        return True


def wrapper_tables(doc):
    """疑似 1×1 外层包裹表：只有“表 X”表题段落 + 一张嵌套表。

    模板外框（1×1，但里面有正文/小节标题/题目表）不算，避免误报。
    返回 `[(标签, 表题文字), ...]`。
    """
    out = []
    for label, tbl in iter_tables(doc.element.body):
        trs = tbl.findall(qn('w:tr'))
        if len(trs) != 1:
            continue
        cells = _row_cells(trs[0])
        if len(cells) != 1 or not _cell_has_table(cells[0]):
            continue
        cell = cells[0]
        paras = cell.findall(qn('w:p'))
        captions = [p for p in paras
                    if re.match(r'^(表|Table)\s*\d', _p_text(p))]
        substantive = [p for p in paras
                       if _p_text(p) and p not in captions]
        if captions and not substantive:
            out.append((label, ' / '.join(_p_text(p) for p in captions)[:80]))
    return out


def _cell_has_table(tc):
    return tc.find(qn('w:tbl')) is not None


def tables_without_blank(doc):
    """顶层表格后没有空段落的表（标签列表）。"""
    out = []
    for label, tbl in iter_tables(doc.element.body):
        if tbl.getparent() is not doc.element.body:
            continue
        nxt = tbl.getnext()
        if nxt is None:
            continue
        if nxt.tag == qn('w:tbl'):
            out.append((label, '表格后紧跟另一张表'))
        elif nxt.tag == qn('w:p') and _p_text(nxt).strip():
            out.append((label, '表格后不是空段落'))
    return out


def table_counts(doc):
    """返回 (顶层表格数, 全部表格数)。用于交付前核对“主表 + N 张数据表”。

    如果全部表格数比预期多出与图/表数量相近的数目，通常就是 1×1 外层表没删干净。
    """
    top = len(tables(doc))
    all_count = len(list(iter_tables(doc.element.body)))
    return top, all_count


def compact_image(src, dst, max_px=1800, colors=256):
    """把截图压成小体积副本再插入，别动原图。

    截图多为 RGBA PNG，直接嵌会让文档暴涨（实测 22 张截图 19.7 MB → 3.1 MB）。缩到
    显示宽度约 300 dpi 再转 256 色，同样的显示尺寸下肉眼无差，体积降 6~7 倍。
    """
    from PIL import Image
    im = Image.open(src).convert('RGB')
    if im.width > max_px:
        im = im.resize((max_px, round(im.height * max_px / im.width)), Image.LANCZOS)
    im.convert('P', palette=Image.ADAPTIVE, colors=colors).save(dst, 'PNG', optimize=True)
    return dst


def compact_images(src_dir, dst_dir, max_px=1800):
    os.makedirs(dst_dir, exist_ok=True)
    out = {}
    for f in sorted(os.listdir(src_dir)):
        if f.lower().endswith(('.png', '.jpg', '.jpeg')):
            out[f] = compact_image(os.path.join(src_dir, f),
                                   os.path.join(dst_dir, f), max_px=max_px)
    return out
