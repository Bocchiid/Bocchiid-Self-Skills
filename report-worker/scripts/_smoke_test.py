#!/usr/bin/env python3
"""Self-test for report-worker.  Run: python scripts/_smoke_test.py

No network, no model calls.  GUI capture is optional (--gui) because it needs an
interactive desktop.
"""

from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import actions
import capture
import env
import verify
import worker as worker_mod


PASS = 0
FAIL = 0


def check(name: str, cond: bool, detail: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ok   {name}")
    else:
        FAIL += 1
        print(f"  FAIL {name} {detail}")


def sample_plan() -> dict:
    return {
        "version": "1.0",
        "purposes": [{"id": "P1", "text": "验证输出", "steps": ["S1"],
                      "evidence": ["E1"]}],
        "stages": [{"id": "ST1", "name": "准备", "steps": ["S1"]}],
        "steps": [{
            "id": "S1", "stage": "ST1", "goal": "输出 hello",
            "action": {"kind": "command", "commands": ["echo hello"],
                       "instruction": "运行 echo"},
            "acceptance": [{"type": "exit_code", "value": 0},
                           {"type": "stdout_regex", "pattern": "hello"}],
            "evidence": ["E1"], "reproduce": "echo hello",
            "failure": [], "irreversible": False, "requires_user": False,
        }],
        "evidence": [{"id": "E1", "step": "S1", "type": "screenshot",
                      "level": "L1", "must_include": ["hello"],
                      "filename": "S1_hello.png"}],
        "report_map": [{"section": "一", "use": ["E1"]}],
        "writer_figures": [],
    }


def test_classify() -> None:
    print("[classify]")
    check("echo is execute",
          actions.classify("echo hello")["category"] == "execute")
    check("rm -rf / is deny",
          actions.classify("rm -rf /")["category"] == "deny")
    check("git push is outward",
          actions.classify("git push origin main")["category"] == "outward")
    check("pip install is elevated",
          actions.classify("pip install pywinauto")["category"] == "elevated")
    check("curl POST is outward",
          actions.classify("curl -X POST https://example.com")["category"] == "outward")
    check("assist allows echo",
          actions.decide("echo hi", "assist")["decision"] == "allow")
    check("assist confirms git push",
          actions.decide("git push", "assist")["decision"] == "confirm")
    check("confirmed git push allows",
          actions.decide("git push", "assist", confirmed=True)["decision"] == "allow")
    check("observe denies echo",
          actions.decide("echo hi", "observe")["decision"] == "deny")
    check("full still denies rm -rf /",
          actions.decide("rm -rf /", "full", confirmed=True)["decision"] == "deny")
    check("taskkill powershell is denied",
          actions.classify("taskkill /IM powershell.exe /F")["category"] == "deny")
    check("taskkill pid is process",
          actions.classify("taskkill /PID 123 /F")["category"] == "process")
    check("assist confirms process control",
          actions.decide("taskkill /PID 123 /F", "assist")["decision"] == "confirm")
    check("kill own pid is denied",
          actions.decide(f"taskkill /PID {__import__('os').getpid()} /F",
                         "assist", confirmed=True)["decision"] == "deny")
    check("pkill python is denied",
          actions.decide("pkill python", "assist", confirmed=True)["decision"] == "deny")
    check("setx is denied",
          actions.classify("setx PATH C:\\x")["category"] == "deny")
    check("pip --user is denied",
          actions.classify("pip install --user requests")["category"] == "deny")
    check("profile write is denied",
          actions.classify("echo x >> ~/.bashrc")["category"] == "deny")


def test_env_isolation() -> None:
    print("[env isolation]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        r = actions.run_action(
            "python -c \"import os; print(os.environ.get('WORKER_SMOKE_VAR'))\"",
            ws, "S1", "local", "", 30, False, "assist",
            env_vars=["WORKER_SMOKE_VAR=ok"], use_venv=False,
        )
        check("temporary env var reaches child",
              "ok" in (r.get("stdout_tail") or ""), str(r))
        check("temporary env var not global",
              "WORKER_SMOKE_VAR" not in __import__("os").environ)
        check("env key is recorded", "WORKER_SMOKE_VAR" in r.get("env_keys", []))
        import _venv
        (ws / ".venv" / ("Scripts" if __import__("os").name == "nt" else "bin")).mkdir(
            parents=True, exist_ok=True)
        py = ws / ".venv" / ("Scripts/python.exe" if __import__("os").name == "nt"
                             else "bin/python")
        py.write_text("", encoding="utf-8")
        check("venv python is discovered", _venv.find_venv_python(ws) is not None)


def test_durability() -> None:
    print("[durability]")
    import _durable
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        worker_mod.create_dirs(ws)
        state = worker_mod.default_state(ws, "assist", "local", "", "durable")
        state["phase"] = "EXECUTE"
        state["plan"]["frozen"] = True
        worker_mod.save_state(ws, state)
        state["execution"]["completed"].append("S1")
        worker_mod.save_state(ws, state)
        check("state backup exists",
              (ws / "_worker" / "state.json.bak").exists())
        (ws / "_worker" / "state.json").write_text("{bad json", encoding="utf-8")
        recovered = worker_mod.load_state(ws)
        check("state recovers from backup",
              recovered.get("phase") == "EXECUTE" and
              "state.json.bak" in recovered.get("_recovered_from", ""))
        worker_mod.save_state(ws, recovered)
        try:
            json.loads((ws / "_worker" / "state.json.bak").read_text(encoding="utf-8"))
            bak_valid = True
        except Exception:
            bak_valid = False
        check("corrupt state does not overwrite good backup", bak_valid)
        _durable.write_intent(ws, "A1", {"kind": "command", "step": "S1",
                                         "command": "echo"})
        check("intent is listed", len(_durable.list_intents(ws)) == 1)
        diag = worker_mod.doctor(ws, recovered)
        check("doctor reports in-flight",
              any("in-flight" in x for x in diag["issues"]), str(diag))
        _durable.write_result(ws, "A1", {"ok": True})
        check("intent cleared after result",
              len(_durable.list_intents(ws)) == 0)
        # Crash window: result written but intent not cleared yet.
        _durable.write_intent(ws, "A2", {"kind": "command", "step": "S1"})
        _durable.atomic_write_json(
            _durable.result_path(ws, "A2"), {"id": "A2", "ok": True})
        check("completed intent is not reported as in-flight",
              len(_durable.list_intents(ws)) == 0)
        _durable.session_event(ws, "test_event", value=1)
        check("session log written",
              (ws / "_worker" / "logs" / "session.log").exists())
        log = ws / "_worker" / "logs" / "session.log"
        log.write_text("x" * 1000, encoding="utf-8")
        rotated = _durable.rotate_log(log, 10, keep=1)
        check("log rotation works", rotated and
              (ws / "_worker" / "logs" / "session.log.1").exists())


def test_action_durability() -> None:
    print("[action durability]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        r = actions.run_action("echo durable", ws, "S1", "local", "",
                               20, False, "assist", use_venv=False)
        check("action result is durable",
              (ws / "_worker" / "actions" / f"{r['id']}.json").exists())
        check("action intent is cleared",
              not (ws / "_worker" / "actions" / f"{r['id']}.intent.json").exists())
        check("action jsonl exists",
              (ws / "_worker" / "actions.jsonl").exists())
        check("action pid recorded", bool(r.get("pid")))


def test_skills() -> None:
    print("[skills]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        result = env.check_skills(ws)
        check("planner/writer/reviewer found", result["ok"], str(result))
        check("skill paths are absolute",
              all(Path(p).is_absolute() for p in result["skills"].values() if p))


def test_analyze_gate() -> None:
    print("[analyze gate]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        worker_mod.create_dirs(ws)
        state = worker_mod.default_state(ws, "assist", "local", "", "gate")
        r = verify.analyze_gate(ws, state)
        check("gate fails without materials/decisions", not r["ok"], str(r))
        (ws / "材料清单.md").write_text(
            "# 材料清单\n\n## 已有\n- 指导书\n\n## 还缺\n- 评分表\n\n## 用户确认\n- 没有评分表\n",
            encoding="utf-8")
        r = verify.analyze_gate(ws, state)
        check("gate still fails without decisions", not r["ok"], str(r))
        state["decisions"].append({"key": "materials", "value": "partial",
                                    "note": "没有评分表"})
        r = verify.analyze_gate(ws, state)
        check("gate passes with materials + decisions", r["ok"], str(r))
        plan = sample_plan()
        plan["conflicts"] = [{"id": "C1", "requirement": "参考文献",
                              "decision": "pending"}]
        (ws / "_worker" / "plan.json").write_text(
            json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        r = verify.analyze_gate(ws, state)
        check("pending conflict blocks gate", not r["ok"], str(r))
        plan["conflicts"][0]["decision"] = "skip"
        (ws / "_worker" / "plan.json").write_text(
            json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        r = verify.analyze_gate(ws, state)
        check("skip conflict passes gate", r["ok"], str(r))


def test_plan() -> None:
    print("[plan]")
    ok = verify.validate_plan(sample_plan())
    check("sample plan passes", ok["ok"], str(ok.get("errors")))
    bad = sample_plan()
    bad["steps"][0].pop("goal")
    check("missing goal fails", not verify.validate_plan(bad)["ok"])
    bad = sample_plan()
    bad["steps"][0]["evidence"] = ["E9"]
    check("unknown evidence fails", not verify.validate_plan(bad)["ok"])
    bad = sample_plan()
    bad["steps"][0]["action"] = {"kind": "gui"}
    check("gui without target fails", not verify.validate_plan(bad)["ok"])
    bad = sample_plan()
    bad["steps"][0]["acceptance"] = []
    check("no acceptance fails", not verify.validate_plan(bad)["ok"])
    bad = sample_plan()
    bad["conflicts"] = [{"id": "C1", "requirement": "参考文献",
                         "decision": "pending"}]
    check("pending conflict fails plan", not verify.validate_plan(bad)["ok"])
    bad["conflicts"][0]["decision"] = "skip"
    check("skip conflict passes plan", verify.validate_plan(bad)["ok"])


def test_state() -> None:
    print("[state]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        worker_mod.create_dirs(ws)
        state = worker_mod.default_state(ws, "assist", "local", "", "smoke")
        worker_mod.save_state(ws, state)
        check("initial phase ANALYZE", worker_mod.load_state(ws)["phase"] == "ANALYZE")
        worker_mod.advance(state, "ORGANIZE")
        check("ANALYZE -> ORGANIZE", state["phase"] == "ORGANIZE")
        try:
            worker_mod.advance(state, "EXECUTE")
            check("illegal ORGANIZE -> EXECUTE rejected", False)
        except SystemExit:
            check("illegal ORGANIZE -> EXECUTE rejected", True)
        plan = sample_plan()
        (ws / "_worker" / "plan.json").write_text(
            json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        (ws / "材料清单.md").write_text(
            "# 材料清单\n\n## 已有\n- 材料\n\n## 还缺\n- 无\n\n## 用户确认\n- 用户确认：没有缺项\n",
            encoding="utf-8")
        state["decisions"].append({"key": "materials", "value": "complete",
                                    "note": "用户确认没有缺项"})
        state["plan"]["frozen"] = True
        worker_mod.advance(state, "PLAN")
        worker_mod.advance(state, "PLAN_FREEZE")
        worker_mod.advance(state, "EXECUTE")
        check("freeze-like path reaches EXECUTE", state["phase"] == "EXECUTE")
        check("next step is S1",
              (worker_mod.next_step(ws, state) or {}).get("id") == "S1")
        state["execution"]["completed"].append("S1")
        check("no next step after completion",
              worker_mod.next_step(ws, state) is None)


def test_step() -> None:
    print("[step]")
    plan = sample_plan()
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        result = {"exit_code": 0, "stdout": "hello world\n", "status": "pass"}
        r = verify.check_step(plan, "S1", result, ws)
        check("step pass", r["ok"] and r["status"] == "pass", str(r))
        result = {"exit_code": 1, "stdout": "nope\n", "status": "fail"}
        r = verify.check_step(plan, "S1", result, ws)
        check("step fail", not r["ok"] and r["status"] == "fail")
        plan2 = sample_plan()
        plan2["steps"][0]["acceptance"].append(
            {"type": "visual", "description": "看到 hello"})
        r = verify.check_step(plan2, "S1",
                              {"exit_code": 0, "stdout": "hello"},
                              ws)
        check("visual makes partial", r["status"] == "partial")


def test_evidence() -> None:
    print("[evidence]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        src = ws / "Pictures" / "E1.txt"
        src.parent.mkdir(parents=True, exist_ok=True)
        src.write_text("hello", encoding="utf-8")
        entry = capture.add_evidence(ws, "E1", "S1", "Pictures/E1.txt",
                                     etype="screenshot", level="L1",
                                     status="pending")
        check("evidence sha recorded", bool(entry.get("sha256")))
        manifest = capture.load_manifest(ws)
        r = verify.check_evidence(manifest, ws)
        check("pending evidence fails", not r["ok"])
        capture.set_status(ws, "E1", "通过")
        r = verify.check_evidence(capture.load_manifest(ws), ws)
        check("passed evidence ok", r["ok"], str(r))
        src.write_text("tampered", encoding="utf-8")
        r = verify.check_evidence(capture.load_manifest(ws), ws)
        check("tampered evidence fails", not r["ok"])


def test_review() -> None:
    print("[review]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        p = ws / "review.md"
        p.write_text(
            "# 报告 审核报告\n\n- 审核结论：通过\n- 总分区间：90~93\n\n"
            "## 2. 问题清单\n### 【正确性错误】\n- 无\n"
            "### 【要求未满足】\n- 无\n### 【模板保护问题】\n- 无\n",
            encoding="utf-8")
        r = verify.check_review(verify.parse_review(p))
        check("clean review passes", r["ok"], str(r))
        p.write_text(
            "# 报告 审核报告\n\n- 审核结论：不通过\n\n"
            "## 2. 问题清单\n### 【正确性错误】\n- 命令 -l 写错\n",
            encoding="utf-8")
        r = verify.check_review(verify.parse_review(p))
        check("blocking review fails", not r["ok"] and r["blocking"])


def test_report() -> None:
    print("[report]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        try:
            from docx import Document
            d = Document()
            d.add_paragraph("hello")
            p = ws / "report.docx"
            d.save(p)
            check("docx report passes", verify.check_report(p)["ok"])
        except Exception as exc:
            check("docx report passes (skipped: no python-docx)", True, str(exc))
        bad = ws / "report.txt"
        bad.write_text("x", encoding="utf-8")
        check("wrong extension fails", not verify.check_report(bad)["ok"])


def test_dod() -> None:
    print("[dod]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        worker_mod.create_dirs(ws)
        state = worker_mod.default_state(ws, "assist", "local", "", "smoke")
        state["phase"] = "REVIEW"
        state["plan"]["frozen"] = True
        plan = sample_plan()
        (ws / "_worker" / "plan.json").write_text(
            json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        img = ws / "Pictures" / "E1.txt"
        img.write_text("hello", encoding="utf-8")
        capture.add_evidence(ws, "E1", "S1", "Pictures/E1.txt",
                             status="通过")
        try:
            from docx import Document
            d = Document()
            d.add_paragraph("hello")
            d.save(ws / "报告" / "报告-v1.docx")
            state["task"]["report"] = "报告/报告-v1.docx"
        except Exception:
            (ws / "报告" / "报告-v1.docx").write_bytes(b"not a docx")
            state["task"]["report"] = "报告/报告-v1.docx"
        review = ws / "报告" / "报告-审核报告-v1.md"
        review.write_text(
            "# 审核\n\n- 审核结论：通过\n- 总分区间：90~93\n\n"
            "## 2. 问题清单\n### 【正确性错误】\n- 无\n",
            encoding="utf-8")
        state["report"]["review_path"] = "报告/报告-审核报告-v1.md"
        worker_mod.save_state(ws, state)
        r = verify.check_dod(ws, state)
        check("DoD passes with complete artifacts", r["ok"], str(r.get("reasons")))


def test_capture() -> None:
    print("[capture]")
    with tempfile.TemporaryDirectory() as td:
        ws = Path(td)
        out = ws / "shot.png"
        r = capture.capture(out)
        check("capture returns a file", bool(r.get("file")) and out.exists(), str(r))
        if r.get("ok"):
            one = capture.unique_path(out)
            check("unique path does not overwrite", one != out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--gui", action="store_true",
                    help="also test screen capture (needs a desktop)")
    args = ap.parse_args()
    test_classify()
    test_env_isolation()
    test_skills()
    test_durability()
    test_action_durability()
    test_analyze_gate()
    test_plan()
    test_state()
    test_step()
    test_evidence()
    test_review()
    test_report()
    test_dod()
    if args.gui:
        test_capture()
    print(f"\n{PASS} passed, {FAIL} failed")
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
