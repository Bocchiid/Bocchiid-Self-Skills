#!/usr/bin/env python3
"""提取 .docx 的文本、格式与结构信息，供实验报告审核使用。

用法：
    python parse_docx.py --check                     # 仅检查环境是否就绪
    python parse_docx.py <文件路径>                   # 解析文本/字体/表格
    python parse_docx.py --audit <报告.docx>          # 结构审计（推荐）
    python parse_docx.py --audit <报告.docx> --template <模板.docx>

`--audit` 检查的是“Word/WPS 跨软件、模板结构、公式、图表”这类肉眼容易漏的问题：
  - compatibilityMode 是否仍是 11（Word 2003，会导致 OMML 公式被转成图片）
  - 表格总数、疑似 1×1 外层包裹表、表格后是否有空段落
  - 模板行数/列数/合并关系/左侧栏目名是否被改动
  - OMML 公式数、InlineShapes 数（保存后公式是否变成图片）
  - 含公式的段落是否显式左对齐、是否关闭 autoSpace
  - 非中文 run 是否缺少显式西文字体（WPS 会把希腊字母/下标回退成宋体）

环境检查输出 JSON：
    {"python": "3.11.0", "python_docx": "1.1.0", "ready": true}
"""

import argparse
import json
import re
import sys

M_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CJK_RE = re.compile(r"[　-〿一-鿿！-｠]")
WESTERN_RE = re.compile(r"[A-Za-zͰ-Ͽ]")


def check_env():
    """检查 Python 与 python-docx 是否可用。"""
    result = {"python": sys.version.split()[0], "python_docx": None, "ready": False}
    try:
        import docx
        result["python_docx"] = getattr(docx, "__version__", "unknown")
        result["ready"] = True
    except ImportError:
        result["install"] = "pip install python-docx"
        result["mirror"] = "pip install python-docx -i https://pypi.tuna.tsinghua.edu.cn/simple"
    return result


def get_font_info(run):
    info = {"text": run.text}
    if run.font.name:
        info["font"] = run.font.name
    if run.font.size:
        info["size_pt"] = run.font.size.pt
    if run.bold:
        info["bold"] = True
    if run.italic:
        info["italic"] = True
    return info


# ---------------------------------------------------------------------------
# 结构审计
# ---------------------------------------------------------------------------

def _qn(tag):
    from docx.oxml.ns import qn
    return qn(tag)


def _text(el):
    return "".join(t.text or "" for t in el.iter(_qn("w:t"))).strip()


def _norm(s):
    return re.sub(r"\s+", "", s or "")


def _effective_cells(tr):
    """去重后的 tc（横向合并时 row.cells 会重复返回同一格）。"""
    return tr.findall(_qn("w:tc"))


def _is_blank_paragraph(p):
    return _norm(_text(p)) == ""


def _cell_has_table(tc):
    return tc.find(_qn("w:tbl")) is not None


def _iter_tables(root, depth=0):
    """递归列出所有 w:tbl，带层级。"""
    for child in root.iterchildren():
        if child.tag == _qn("w:tbl"):
            yield depth, child
            yield from _iter_tables(child, depth + 1)
        else:
            yield from _iter_tables(child, depth)


def _table_shape(tbl):
    rows = []
    for tr in tbl.findall(_qn("w:tr")):
        cells = []
        for tc in _effective_cells(tr):
            tcPr = tc.find(_qn("w:tcPr"))
            span, vmerge = 1, ""
            if tcPr is not None:
                gs = tcPr.find(_qn("w:gridSpan"))
                if gs is not None:
                    try:
                        span = int(gs.get(_qn("w:val"), "1"))
                    except ValueError:
                        pass
                vm = tcPr.find(_qn("w:vMerge"))
                if vm is not None:
                    vmerge = vm.get(_qn("w:val"), "continue")
            cells.append({"span": span, "vmerge": vmerge,
                          "label": _norm(_text(tc))})
        rows.append(cells)
    return {"rows": len(rows), "cells": rows}


def _compat_mode(doc):
    try:
        settings = doc.settings.element
    except Exception:
        return None
    compat = settings.find(_qn("w:compat"))
    if compat is None:
        return None
    # 现代 Word：<w:compatSetting w:name="compatibilityMode" w:val="15"/>
    for node in compat.findall(_qn("w:compatSetting")):
        if node.get(_qn("w:name")) == "compatibilityMode":
            try:
                return int(node.get(_qn("w:val"), "0"))
            except ValueError:
                return None
    # 旧写法：<w:compatibilityMode w:val="15"/>
    node = compat.find(_qn("w:compatibilityMode"))
    if node is None:
        return None
    try:
        return int(node.get(_qn("w:val"), "0"))
    except ValueError:
        return None


def _para_omml_count(p):
    return len(list(p.iter(M_NS + "oMath"))) + len(list(p.iter(M_NS + "oMathPara")))


def _math_text(el):
    return "".join(t.text or "" for t in el.iter(M_NS + "t")).strip()


def _audit_omml_structure(doc):
    """检查 OMML 的常见结构错误。

    - 分式 `m:den` / `m:num` 为空，或分母跑到分式外；
    - 上下标 `m:e` / `m:sub` / `m:sup` 为空占位；
    - 求和上下限 `m:lim` 为空；
    - 函数 `m:fName` / `m:e` 为空；
    - 函数参数疑似被并进 `m:sub`（如 `Rα(T)` 的 `(T)`）。
    """
    issues = []

    def need(parent, tag, label):
        node = parent.find(M_NS + tag)
        if node is None:
            issues.append({"structure": label, "issue": f"缺少 {tag}",
                           "text": _math_text(parent)[:80]})
            return False
        if not _math_text(node):
            issues.append({"structure": label, "issue": f"{tag} 为空",
                           "text": _math_text(parent)[:80]})
            return False
        return True

    for f in doc.element.body.iter(M_NS + "f"):
        need(f, "num", "分数")
        need(f, "den", "分数")
    for s in doc.element.body.iter(M_NS + "sSub"):
        need(s, "e", "下标")
        need(s, "sub", "下标")
        sub = s.find(M_NS + "sub")
        if sub is not None:
            sub_text = _math_text(sub)
            if sub_text.startswith("(") and sub_text.endswith(")"):
                issues.append({
                    "structure": "下标",
                    "issue": "函数参数疑似落在 m:sub 里",
                    "text": _math_text(s)[:80],
                    "sub": sub_text[:40],
                })
    for s in doc.element.body.iter(M_NS + "sSup"):
        need(s, "e", "上标")
        need(s, "sup", "上标")
    for s in doc.element.body.iter(M_NS + "sSubSup"):
        need(s, "e", "上下标")
        need(s, "sub", "上下标")
        need(s, "sup", "上下标")
    for n in doc.element.body.iter(M_NS + "nary"):
        need(n, "e", "求和")
        for tag in ("sub", "sup"):
            node = n.find(M_NS + tag)
            if node is not None and not _math_text(node):
                issues.append({"structure": "求和", "issue": f"{tag} 为空占位",
                               "text": _math_text(n)[:80]})
    for fn in doc.element.body.iter(M_NS + "func"):
        need(fn, "fName", "函数")
        need(fn, "e", "函数")
    for lim in doc.element.body.iter(M_NS + "limLow"):
        need(lim, "e", "下限")
        need(lim, "lim", "下限")
    for lim in doc.element.body.iter(M_NS + "limUpp"):
        need(lim, "e", "上限")
        need(lim, "lim", "上限")
    return issues


def _audit_formula_paragraphs(doc):
    """含 OMML 的行内/显示公式段落：对齐与 autoSpace。"""
    problems = []
    for p in doc.element.body.iter(_qn("w:p")):
        if not _para_omml_count(p):
            continue
        pPr = p.find(_qn("w:pPr"))
        jc = pPr.find(_qn("w:jc")) if pPr is not None else None
        alignment = jc.get(_qn("w:val")) if jc is not None else None
        if alignment not in ("left", "start"):
            problems.append({
                "issue": "公式段落没有显式左对齐",
                "alignment": alignment or "(继承样式)",
                "text": _text(p)[:60],
                "fix": "设置 w:jc=left；若继承 Normal 的两端对齐，公式会被拉伸。",
            })
        for tag, name in (("w:autoSpaceDE", "autoSpaceDE"),
                          ("w:autoSpaceDN", "autoSpaceDN")):
            node = pPr.find(_qn(tag)) if pPr is not None else None
            val = node.get(_qn("w:val")) if node is not None else None
            if val not in ("0", "false", "off"):
                problems.append({
                    "issue": f"公式段落未关闭 {name}",
                    "value": val or "(继承)",
                    "text": _text(p)[:60],
                    "fix": "显式写 <w:autoSpaceDE w:val=\"0\"/> / autoSpaceDN，避免 WPS 角标被拉开。",
                })
    return problems


def _audit_run_fonts(doc):
    """非中文 run 是否缺显式西文字体。WPS 会回退到宋体，按中文字宽排。"""
    missing = []
    total = 0
    for r in doc.element.body.iter(_qn("w:r")):
        text = "".join(t.text or "" for t in r.iter(_qn("w:t")))
        if not text or not WESTERN_RE.search(text) or CJK_RE.search(text):
            continue
        total += 1
        rPr = r.find(_qn("w:rPr"))
        fonts = rPr.find(_qn("w:rFonts")) if rPr is not None else None
        ascii_font = fonts.get(_qn("w:ascii")) if fonts is not None else None
        if not ascii_font:
            missing.append(text[:40])
    return total, missing


def _audit_tables(doc):
    all_tables = list(_iter_tables(doc.element.body))
    top_tables = [t for depth, t in all_tables if depth == 0 and
                  t.getparent() is doc.element.body]
    wrappers = []
    no_blank_after = []
    for idx, (depth, tbl) in enumerate(all_tables):
        trs = tbl.findall(_qn("w:tr"))
        if len(trs) == 1:
            cells = _effective_cells(trs[0])
            if len(cells) == 1:
                cell = cells[0]
                paras = cell.findall(_qn("w:p"))
                caption_re = re.compile(r"^(表|Table)\s*\d")
                captions = [p for p in paras
                            if caption_re.match(_text(p))]
                # 只有“表题段落 + 嵌套表、没有别的正文”才像 1×1 外层包裹表。
                # 模板外框（1×1，里面还有正文/小节标题/题目表）不应误报。
                substantive = [p for p in paras
                               if _text(p) and p not in captions]
                if _cell_has_table(cell) and captions and not substantive:
                    wrappers.append({
                        "table_index": idx,
                        "text": " / ".join(_text(p) for p in captions)[:80],
                        "reason": "1×1 表格只包了“表题+表格”：WPS 会显示成表格外多一圈框。",
                    })
                elif _cell_has_table(cell) and len(paras) == 0:
                    wrappers.append({
                        "table_index": idx,
                        "text": "1×1 表格里只有一张嵌套表",
                        "reason": "疑似 1×1 外层包裹表：WPS 会显示成表格外多一圈框。",
                    })
    for tbl in top_tables:
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == _qn("w:tbl"):
            no_blank_after.append({"issue": "表格后紧跟另一张表，没有空行"})
        elif nxt is not None and nxt.tag == _qn("w:p") and not _is_blank_paragraph(nxt):
            no_blank_after.append({
                "issue": "表格后没有空段落",
                "next_text": _text(nxt)[:60],
            })
    return {
        "top_level": len(top_tables),
        "all": len(all_tables),
        "nested": len(all_tables) - len(top_tables),
        "wrappers": wrappers,
        "no_blank_after": no_blank_after,
    }


def _shape_summary(doc):
    return [_table_shape(tbl) for _, tbl in _iter_tables(doc.element.body)]


def _row_pattern(shape):
    return tuple(len(row) for row in shape["cells"])


def _compare_template(report_doc, template_doc):
    """模板行数/列数/合并/栏目名是否与报告一致。

    按“表格形状”匹配，不按文档顺序：writer 会在模板单元格里插入数据表，
    顺序比对会把后面的题目表错位成数据表，产生假阳性。
    """
    tpl = _shape_summary(template_doc)
    rep = _shape_summary(report_doc)
    issues = []
    used = set()

    def label_score(a, b):
        score = 0
        for ra, rb in zip(a["cells"], b["cells"]):
            for ca, cb in zip(ra, rb):
                if ca["label"] and ca["label"] == cb["label"]:
                    score += 1
        return score

    def find_match(shape, same_rows=True):
        best, best_score = None, -1
        for j, cand in enumerate(rep):
            if j in used:
                continue
            if same_rows:
                if _row_pattern(cand) != _row_pattern(shape):
                    continue
                if cand["rows"] != shape["rows"]:
                    continue
            else:
                # 行数不同时只要求列结构一致（首行列数 + 每行列数模式放宽）
                if (not cand["cells"] or not shape["cells"]
                        or len(cand["cells"][0]) != len(shape["cells"][0])):
                    continue
                if cand["rows"] == shape["rows"]:
                    continue
            score = label_score(shape, cand)
            # 模板有固定栏目名时，至少要对上一个，避免把数据表误认成模板表
            if any(c["label"] for row in shape["cells"] for c in row) and score == 0:
                continue
            if score > best_score:
                best, best_score = j, score
        return best

    for i, a in enumerate(tpl):
        j = find_match(a, same_rows=True)
        if j is None:
            # 列结构相同但行数不同 → 典型的“加行/删行”
            j = find_match(a, same_rows=False)
            if j is not None:
                used.add(j)
                issues.append({
                    "table": i, "issue": "行数变了",
                    "template": a["rows"], "report": rep[j]["rows"],
                    "fix": "模板行数必须保持原样；摘要/结论应嵌入现有栏目，不要加行。",
                })
            else:
                issues.append({
                    "table": i, "issue": "模板表格缺失或被替换",
                    "rows": a["rows"],
                    "fix": "模板表格必须保留；不要删除/重建模板结构。",
                })
            continue
        used.add(j)
        b = rep[j]
        for ri, (ra, rb) in enumerate(zip(a["cells"], b["cells"])):
            if len(ra) != len(rb):
                issues.append({
                    "table": i, "row": ri, "issue": "列数变了",
                    "template": len(ra), "report": len(rb),
                })
                continue
            for ci, (ca, cb) in enumerate(zip(ra, rb)):
                if ca["span"] != cb["span"] or ca["vmerge"] != cb["vmerge"]:
                    issues.append({
                        "table": i, "row": ri, "col": ci,
                        "issue": "合并关系变了",
                        "template": [ca["span"], ca["vmerge"]],
                        "report": [cb["span"], cb["vmerge"]],
                    })
                # 空栏目名是填写区，允许被填内容；只比较模板原本有字的栏目名
                if ca["label"] and ca["label"] != cb["label"]:
                    issues.append({
                        "table": i, "row": ri, "col": ci,
                        "issue": "左侧栏目名/固定文字变了",
                        "template": ca["label"][:40],
                        "report": cb["label"][:40],
                    })
    extra = len(rep) - len(used)
    if extra > 0:
        issues.append({
            "issue": f"报告比模板多出 {extra} 张表",
            "note": "数据表/附录表可能是正常的；若多出的数量接近图表数，先查 1×1 外层包裹表。",
        })
    return issues


def audit(path, template=None):
    from docx import Document

    doc = Document(path)
    report = {
        "file": str(path),
        "compatibility_mode": _compat_mode(doc),
        "omml_count": len(list(doc.element.body.iter(M_NS + "oMath"))),
        "inline_shapes": len(doc.inline_shapes),
        "tables": _audit_tables(doc),
        "formula_paragraphs": _audit_formula_paragraphs(doc),
        "omml_structure": _audit_omml_structure(doc),
    }
    total_runs, missing_fonts = _audit_run_fonts(doc)
    report["western_runs"] = {
        "total": total_runs,
        "missing_explicit_font": len(missing_fonts),
        "samples": missing_fonts[:5],
    }
    if report["compatibility_mode"] not in (None, 15):
        report.setdefault("warnings", []).append(
            "compatibilityMode=%s（不是 15）：Word 保存时可能把 OMML 公式转成图片。"
            % report["compatibility_mode"])
    if report["tables"]["wrappers"]:
        report.setdefault("warnings", []).append(
            "发现 %d 处疑似 1×1 外层包裹表；WPS 会显示多余外框。"
            % len(report["tables"]["wrappers"]))
    if report["tables"]["no_blank_after"]:
        report.setdefault("warnings", []).append(
            "有表格后没有空段落；不满足“表格后空一行”。")
    if report["omml_structure"]:
        report.setdefault("warnings", []).append(
            "发现 %d 处 OMML 结构问题（空分母/空上下标/函数参数进下标）。"
            % len(report["omml_structure"]))
    if template:
        report["template_compare"] = _compare_template(doc, Document(template))
    return report


def parse(path):
    from docx import Document

    doc = Document(path)
    result = {"paragraphs": []}

    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        entry = {
            "text": para.text,
            "style": para.style.name if para.style else None,
            "runs": [get_font_info(r) for r in para.runs if r.text],
        }
        if para.paragraph_format.alignment is not None:
            entry["alignment"] = str(para.paragraph_format.alignment)
        result["paragraphs"].append(entry)

    result["tables"] = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        result["tables"].append(rows)

    result["audit"] = audit(path)
    return result


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path", nargs="?", help="docx path")
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--audit", metavar="REPORT", help="audit report structure")
    ap.add_argument("--template", help="compare report tables against this template")
    args = ap.parse_args()

    if args.check:
        print(json.dumps(check_env(), ensure_ascii=False, indent=2))
        return 0

    env = check_env()
    if not env["ready"]:
        print(json.dumps({
            "error": "python-docx 未安装",
            "install": env.get("install"),
            "mirror": env.get("mirror"),
        }, ensure_ascii=False, indent=2))
        return 1

    if args.audit:
        print(json.dumps(audit(args.audit, args.template),
                         ensure_ascii=False, indent=2))
        return 0

    if not args.path:
        print("用法：python parse_docx.py --check | <文件路径> | "
              "--audit <报告.docx> [--template <模板.docx>]")
        return 1
    print(json.dumps(parse(args.path), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
