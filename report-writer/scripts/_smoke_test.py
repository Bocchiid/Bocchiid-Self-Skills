# -*- coding: utf-8 -*-
"""docx_kit / diagram_kit 的冒烟测试。

构造一份"像真模板"的文档：封面表单 + 外框 + 框内嵌的题目表 + 正文填写区，
然后跑一遍：勘察 → 填封面 → 正文写作 → 模板骨架校验；
再生成一张示例架构图并跑几何校验，确认能抓到节点重叠/箭头缺失。

    python scripts/_smoke_test.py
"""
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from docx import Document          # noqa: E402
import docx_kit as K               # noqa: E402

OK = [0]
BAD = [0]


def check(name, cond, extra=''):
    print('%s %s%s' % ('  ok ' if cond else '  XX ', name,
                       ('  → %s' % extra) if extra and not cond else ''))
    (OK if cond else BAD)[0] += 1


def build_template(path):
    d = Document()
    d.add_paragraph('课程名称：示例课程')          # 段落式封面（已填）
    d.add_paragraph('题　目：')                      # 段落式封面（空，无下划线）
    d.add_paragraph('')

    # ① 封面表单：标签格 | 值格
    t = d.add_table(rows=4, cols=4)
    t.style = 'Table Grid'
    t.cell(0, 0).text = '实验名称：'
    t.cell(0, 1).text = '示例课程一'             # 已填
    t.cell(1, 0).text = '姓　　名：'
    t.cell(1, 1).text = ''
    t.cell(1, 2).text = '学　　号：'
    t.cell(1, 3).text = ''
    t.cell(2, 0).text = '班　　级：'
    t.cell(2, 1).text = ''
    t.cell(2, 2).text = '日　　期：'
    t.cell(2, 3).text = ''
    t.cell(3, 0).text = '题　　目：'
    t.cell(3, 1).text = '＿＿＿＿＿＿'               # 下划线空格（模板留的空白）

    # ② 外框（单列长表），框里嵌一张题目表
    frame = d.add_table(rows=3, cols=1)
    frame.style = 'Table Grid'
    frame.cell(0, 0).text = '一、实验目的：1、掌握 vim 的使用'
    inner_label = frame.cell(1, 0).text = '三、题目信息'   # 先写一行标题
    inner = frame.cell(1, 0).add_table(rows=2, cols=2)     # 内嵌表格
    inner.style = 'Table Grid'
    inner.cell(0, 0).text = '题　目'
    inner.cell(0, 1).text = ''
    inner.cell(1, 0).text = '指导教师'
    inner.cell(1, 1).text = ''
    frame.cell(2, 0).text = '四、基础实验步骤与结果：'
    d.save(path)
    return inner_label


def main():
    tmp = tempfile.mkdtemp(prefix='docxkit-smoke-')
    tpl = os.path.join(tmp, '模板.docx')
    build_template(tpl)

    print('\n=== 1. 勘察 ===')
    d = Document(tpl)
    s = K.survey(d)
    check('survey 标出内嵌表格', '内嵌表' in s)
    check('survey 统计到 1 张内嵌表格',
          K.survey(d).count('内嵌表[') == 1,
          [l for l in s.split('\n') if '内嵌' in l])
    check('nested_tables 找到 1 张', len(K.nested_tables(d)) == 1,
          K.nested_tables(d))

    print('\n=== 2. 封面字段勘察（只报告，不改动）===')
    before_xml = d.element.body.xml
    fields = K.cover_fields(d)
    for f in fields:
        print('    %-8s %-14s %-24s %s' % (f['kind'], f['label'],
                                           f['value'][:16] or '(空)', f['where']))
    check('勘察封面不碰文档', before_xml == d.element.body.xml)
    labels = {f['label']: f for f in fields}
    for want in ('姓名', '学号', '班级', '日期', '题目', '课程名称', '实验名称',
                 '指导教师'):
        check('识别出「%s」' % want, want in labels)
    check('无冒号的题目表标签 → next_cell',
          labels.get('题目', {}).get('kind') in ('next_cell', 'inline'))
    check('已填字段的值读出来了',
          labels.get('实验名称', {}).get('value') == '示例课程一',
          labels.get('实验名称'))
    check('下划线不算已填的值', labels.get('题目', {}).get('value') == '',
          labels.get('题目'))

    print('\n=== 3. 填封面：标签、下划线、内嵌表都不能动 ===')
    before = K.snapshot(d)
    filled, missed, manual = K.fill_cover(d, {
        '姓名': '示例姓名', '学号': '00000000', '班级': '示例班级',
        '日期': '2026年10月16日', '题目': '示例课程实验',
        '指导教师': '示例教师'})
    for line in filled:
        print('    填 ', line)
    check('六项全部命中', not missed, missed)
    check('没有需要人工处理的字段', not manual, manual)
    struct = K.diff_structure(before, K.snapshot(d))
    check('结构零变化（行数/边框/表格数）', not struct, struct)
    txt = d.element.body.xml
    check('标签「姓　　名：」还在', '姓' in txt and '名' in txt)
    check('内嵌表格还在', len(K.nested_tables(d)) == 1)
    check('下划线空格被值替换', '＿＿＿＿＿＿' not in txt)
    inner = K.nested_tables(d)[0][1]
    check('内嵌题目表填上了值',
          K._tc_text(K._row_cells(inner.findall(K.qn('w:tr'))[0])[1])
          == '示例课程实验',
          K._tc_text(K._row_cells(inner.findall(K.qn('w:tr'))[0])[1]))
    check('内嵌题目标签格没被改',
          K._tc_text(K._row_cells(inner.findall(K.qn('w:tr'))[0])[0]) == '题　目')
    # 覆盖值：已填的格子再填一次，不应变成拼接
    K.fill_cover(d, {'实验名称': '示例课程二'})
    check('已填字段被替换而不是拼接',
          '示例课程二' in d.element.body.xml and '示例课程一' not in d.element.body.xml)

    print('\n=== 4. 写入前拦截 ===')
    d2 = Document(tpl)
    print('  -- 锚点落在封面字段上 --')
    try:
        cur = K.Cursor(d2, K.Target.after('题　目：'))
        cur.put(K.heading(1, '1. 往封面后面插正文'))
        check('封面字段锚点被拦下', False, '没拦到')
    except ValueError as e:
        check('封面字段锚点被拦下', '封面' in str(e), str(e)[:60])
    print('  -- 正文锚点不误伤 --')
    cur = K.Cursor(d2, K.Target.after('一、实验目的：1、掌握 vim 的使用'))
    cur.put(K.body(K.Captions(), '补充说明。'))
    rep = cur.finish()
    check('正文锚点正常放行', not cur.check() and all('row0' in r for r in rep),
          (cur.check(), rep))
    print('  -- 封面行 --')
    try:
        cur = K.Cursor(d2, K.Target.cell(table=0, row=1), keep_paras=0)
        cur.put(K.heading(1, '1. 往封面里写'))
        check('封面行被拦下', False, '没拦到')
    except ValueError as e:
        check('封面行被拦下', '封面' in str(e), str(e)[:60])
    print('  -- 含内嵌表格的格子 --')
    try:
        cur = K.Cursor(d2, K.Target.cell(table=1, row=1), keep_paras=0)
        cur.put(K.heading(1, '1. 清空含内嵌表的格子'))
        check('内嵌表格所在格被拦下', False, '没拦到')
    except ValueError as e:
        check('内嵌表格所在格被拦下', '内嵌' in str(e), str(e)[:80])
    print('  -- 给理由可放行（clear=False 追加）--')
    cur = K.Cursor(d2, K.Target.cell(table=1, row=1),
                   keep_paras=0, clear=False)
    cur.force('人工确认：只往内嵌表格后面追加一段，不清空这一格')
    cur.put(K.body(K.Captions(), '补充说明。'))
    rep = cur.finish()
    check('clear=False 追加后内嵌表仍在、行数没变',
          len(K.nested_tables(d2)) == 1 and not any('行数' in r for r in rep), rep)

    print('\n=== 5. 正常写正文（单列长表最后一格）===')
    d3 = Document(tpl)
    caps = K.Captions()
    caps.declare('fig1', '图')
    cur = K.Cursor(d3, K.Target.cell(table=1, row=2), keep_paras=1)
    cur.put(K.heading(1, '1. 环境准备'))
    cur.put(K.body(caps, '详见<<fig1>>。'))
    rep = cur.finish()
    check('只有目标格有字数变化',
          len(rep) == 1 and 'row2' in rep[0], rep)
    check('内嵌表没被动', len(K.nested_tables(d3)) == 1)

    out = os.path.join(tmp, 'out.docx')
    d3.save(out)
    ok = Document(out)
    check('产出能重新打开', len(K.nested_tables(ok)) == 1)

    print('\n=== 6. --plan / --cover 命令 ===')
    import subprocess
    env = dict(os.environ, PYTHONIOENCODING='utf-8')
    for args in (['--plan', tpl], ['--cover', tpl]):
        r = subprocess.run([sys.executable, 'scripts/docx_env.py'] + args,
                           cwd=os.path.dirname(os.path.dirname(
                               os.path.abspath(__file__))),
                           capture_output=True, text=True, encoding='utf-8',
                           env=env)
        check('docx_env.py %s 正常输出' % args[0],
              r.returncode == 0 and len(r.stdout) > 50, r.stderr[-200:])
    r = subprocess.run([sys.executable, 'scripts/docx_env.py', '--cover', tpl],
                       cwd=os.path.dirname(os.path.dirname(
                           os.path.abspath(__file__))),
                       capture_output=True, text=True, encoding='utf-8', env=env)
    print(r.stdout)
    check('--cover 列出了待填字段', '待填' in r.stdout)
    check('--cover 提示别用 put() 写封面', 'put()' in r.stdout)

    print('\n=== 7. run 级格式保真（封面最容易被写坏的地方）===')
    p2 = os.path.join(tmp, 'runs.docx')
    d4 = Document()
    d4.add_paragraph('占位')
    t2 = d4.add_table(rows=2, cols=2)
    t2.style = 'Table Grid'
    # 逐字拆 run + 两段下划线空白：模板常见的"书法格子"写法
    lab = t2.cell(0, 0).paragraphs[0]
    for ch in ('姓', '　', '名', '：'):
        lab.add_run(ch)
    for _ in range(2):
        r = lab.add_run('＿')
        r.underline = True
    t2.cell(1, 0).paragraphs[0].add_run('题　目：示例实验')   # 标签和值挤在一个 run
    d4.save(p2)

    d5 = Document(p2)
    pl = d5.tables[0].cell(0, 0).paragraphs[0]
    reps_before = len(K._tail_runs(pl, '姓名'))
    check('标签后精确取到 2 个空白 run', reps_before == 2, reps_before)
    base = K.snapshot(d5)
    K.fill_cover(d5, {'姓名': '示例姓名', '题目': '新题目'})
    check('结构没变', not K.diff_structure(base, K.snapshot(d5)),
          K.diff_structure(base, K.snapshot(d5)))
    pl = d5.tables[0].cell(0, 0).paragraphs[0]
    check('逐字拆开的标签没被动', K._p_text(pl).startswith('姓　名：'),
          K._p_text(pl))
    check('值是「示例姓名」', K._p_text(pl) == '姓　名：示例姓名', K._p_text(pl))
    runs = pl._p.findall(K.qn('w:r'))
    filled = [r for r in runs if K._run_text(r) == '示例姓名']
    check('值写进了原来那个下划线 run', len(filled) == 1)
    check('下划线保留', filled and filled[0].find(K.qn('w:rPr')) is not None
          and K._run_text(filled[0]) == '示例姓名'
          and filled[0].find(K.qn('w:rPr')).find(K.qn('w:u')) is not None,
          K._run_text(filled[0]))
    check('多余的空白下划线被清空',
          sum(1 for r in runs if K._run_text(r) and
              K._run_text(r).strip('_＿') == '') == 0,
          [K._run_text(r) for r in runs])
    p22 = d5.tables[0].cell(1, 0).paragraphs[0]
    check('一个 run 里的「标签：值」被正确替换',
          K._p_text(p22) == '题　目：新题目', K._p_text(p22))
    check('替换没有把值拆成两个 run',
          len([r for r in p22._p.findall(K.qn('w:r')) if '新题目' in K._run_text(r)]) == 1)

    # w:u 的插入位置必须合法：sz 之后、vertAlign 之前
    prop = K.rpr(sz=21, superscript=False, underline=True)
    order = [c.tag.split('}')[-1] for c in prop]
    check('rpr 子元素顺序合法', order == ['rFonts', 'sz', 'u'], order)
    sup = K.rpr(sz=21, superscript=True, underline=True)
    order2 = [c.tag.split('}')[-1] for c in sup]
    check('u 排在 vertAlign 之前', order2 == ['rFonts', 'sz', 'u', 'vertAlign'],
          order2)
    d6 = Document()
    para = d6.add_paragraph()
    r = para.add_run('x')
    r.font.size = None
    r._r.find(K.qn('w:rPr')).append(K.E('w:vertAlign', val='superscript'))
    K._ensure_underline(r._r)
    order3 = [c.tag.split('}')[-1] for c in r._r.find(K.qn('w:rPr'))]
    check('已有 vertAlign 时 u 插到它前面',
          order3.index('u') < order3.index('vertAlign'), order3)

    print('\n=== 8. 往含题目表格的单元格里写内容（推荐路径：Target.after）===')
    d7 = Document(tpl)
    inner_before = K.nested_tables(d7)[0][1]
    prof_before = K._table_profile(inner_before)
    cur = K.Cursor(d7, K.Target.after('三、题目信息'))
    cur.put(K.body(K.Captions(), '补充：本实验的题目信息见下表。'))
    rep = cur.finish()
    inner_after = K.nested_tables(d7)[0][1]
    check('内嵌题目表一个字没动（行列/字数/边框）',
          K._table_profile(inner_after) == prof_before,
          (prof_before, K._table_profile(inner_after)))
    check('finish 只报所在格的文字变化',
          rep and all('内嵌表' not in r for r in rep), rep)
    tc = inner_after.getparent()
    kids = [c.tag.split('}')[-1] for c in tc.iterchildren()
            if c.tag in (K.qn('w:p'), K.qn('w:tbl'))]
    check('新段落插在标签段之后、题目表之前', kids[:3] == ['p', 'p', 'tbl'], kids)

    # 往单元格里插表格（锚点是段落）→ 必须补一个 1pt 段落，
    # 否则单元格以 w:tbl 结尾，Word 打开时报"文件内容已经损坏"
    d8 = Document(tpl)
    caps8 = K.Captions()
    caps8.declare('t9', '表')
    cur = K.Cursor(d8, K.Target.after('四、基础实验步骤与结果：'))
    cur.put(K.wrapped_table(caps8, 't9', '测试表',
                            [['a', 'b'], ['1', '2']], [2000, 2000]))
    cur.finish()
    frame_cell = K._row_cells(K.tables(d8)[1].findall(K.qn('w:tr'))[2])[0]
    kids = [c.tag.split('}')[-1] for c in frame_cell.iterchildren()
            if c.tag in (K.qn('w:p'), K.qn('w:tbl'))]
    check('单元格以段落收尾（Word 不会报损坏）',
          kids[-1] == 'p' and kids[-2] == 'tbl', kids)

    # table_block 不能再生成 1×1 外层表；表后必须有空段
    d9 = Document()
    caps9 = K.Captions()
    caps9.declare('t10', '表')
    cur = K.Cursor(d9, K.Target.at_end())
    cur.put_table(caps9, 't10', '测试表', [['a', 'b'], ['1', '2']], [2000, 2000])
    check('table_block 只生成一张表（没有 1×1 外层表）',
          K.table_counts(d9) == (1, 1), K.table_counts(d9))
    last_two = [c.tag.split('}')[-1] for c in d9.element.body.iterchildren()
                if c.tag in (K.qn('w:p'), K.qn('w:tbl'))][-2:]
    check('表格后有空段落', last_two == ['tbl', 'p'], last_two)

    # 模板骨架能抓到“加行”
    shape_before = K.template_shape(d9)
    d9.tables[0].add_row()
    check('template_shape 抓到行数变化',
          bool(K.diff_template_shape(shape_before, K.template_shape(d9))))

    # 跨文档比对（交付前 verify_output 用）：重新打开文件也不能误报
    tpl_path = os.path.join(tmp, 'shape-tpl.docx')
    rep_path = os.path.join(tmp, 'shape-rep.docx')
    dt = Document()
    tt = dt.add_table(rows=2, cols=2)
    tt.cell(0, 0).text = '栏目A'
    tt.cell(0, 1).text = '值'
    tt.cell(1, 0).text = '栏目B'
    dt.save(tpl_path)
    dr = Document(tpl_path)
    dr.tables[0].add_row().cells[0].text = '摘要'
    dr.save(rep_path)
    issues = K.diff_template_documents(Document(tpl_path), Document(rep_path))
    check('跨文档骨架比对标出加行',
          any('行数变了' in x for x in issues), issues)
    same = K.diff_template_documents(Document(tpl_path), Document(tpl_path))
    check('跨文档骨架比对同文件不误报', not same, same)

    # 示意图：生成 → 几何校验 → 故意造错能报出来
    print('\n=== 9. diagram_kit / check_diagram ===')
    import diagram_kit as D
    import check_diagram as C
    spec = D.example_spec()
    layout = D.compute_layout(spec)
    svg = os.path.join(tmp, 'diagram.svg')
    png = os.path.join(tmp, 'diagram.png')
    D.render(layout, __import__('pathlib').Path(svg),
             __import__('pathlib').Path(png), dpi=200)
    lay_path = os.path.join(tmp, 'diagram.layout.json')
    with open(lay_path, 'w', encoding='utf-8') as f:
        json.dump(layout, f, ensure_ascii=False)
    res = C.check(layout)
    check('生成的架构图通过几何校验', res['ok'], res)
    check('没有节点重叠/穿框/箭头缺失',
          not res['errors'], res['errors'])
    bad = json.loads(json.dumps(layout))
    bad['nodes'][1]['rect']['x'] = bad['nodes'][0]['rect']['x']
    bad['nodes'][1]['rect']['y'] = bad['nodes'][0]['rect']['y']
    check('几何校验能抓到节点重叠', not C.check(bad)['ok'])
    bad2 = json.loads(json.dumps(layout))
    bad2['edges'][0].pop('arrow', None)
    check('几何校验能抓到箭头缺失', not C.check(bad2)['ok'])

    print('\n通过 %d 项，失败 %d 项' % (OK[0], BAD[0]))
    print('临时目录：%s' % tmp)
    return 1 if BAD[0] else 0


if __name__ == '__main__':
    sys.exit(main())
